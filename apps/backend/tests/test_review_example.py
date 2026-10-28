"""The explicit sample import exercises real review rules without calling a provider."""

from dataclasses import asdict
from io import BytesIO
from types import SimpleNamespace

from openpyxl import load_workbook

from app.jobs.seed_review_example import build_review_example
from app.review.engine import project, synchronize
from app.review.service import State, evaluate
from app.review.taxonomy import CATEGORIES
from tests.test_grouped_review import operation


def example_state():
    document, source = build_review_example()
    request = State(
        request_id="example-test",
        workflow_status="review",
        review_revision=0,
        selected_sheet=document.selected_sheet,
        selected_sheets=document.selected_sheets,
        parser_warnings=document.warnings,
        ai_review={"sources": {}},
        items=[],
        review_issues=[],
    )
    for index, parsed in enumerate(document.items, 1):
        request.items.append(
            SimpleNamespace(
                **asdict(parsed),
                id=index,
                field_reviews={},
                source_active=True,
                removed_at=None,
                removed_by=None,
                source_reference=SimpleNamespace(page=0, row=parsed.row, excerpt=parsed.excerpt),
            )
        )
        request.ai_review["sources"][str(index)] = parsed.review_source
    return request, source


def test_example_covers_all_problem_codes_and_has_real_source_evidence():
    request, raw = example_state()
    projection = synchronize(request)
    assert {i["code"] for i in projection["issues"]} == set(CATEGORIES)
    assert {t["category"] for t in projection["tasks"]} == set(CATEGORIES.values())
    mapping = [i for i in projection["issues"] if i.get("scope") == "table"]
    assert len(mapping) == 1
    assert len(mapping[0]["itemIds"]) == 2
    assert projection["readyIds"] == [1, 2, len(request.items)]
    workbook = load_workbook(BytesIO(raw))
    assert set(workbook.sheetnames) == {"Review cases", "Column check"}
    for item in request.items:
        source = request.ai_review["sources"][str(item.id)]
        assert workbook[source["worksheet"]].cell(item.source_reference.row, 1).value == (
            source["cells"]["0"] or None
        )
    assert "simulated" in request.parser_warnings[0]


def test_example_overlap_requires_both_decisions_and_column_confirmation_is_shared():
    request, _ = example_state()
    projection = synchronize(request)
    overlapping = next(
        i for i in request.items if i.attributes["Example case"] == "Overlapping problems"
    )
    corrected, _, _, _ = evaluate(
        request, operation(request, "edit", [overlapping.id], changes={"domain": "equipment"})
    )
    assert overlapping.id not in project(corrected)["readyIds"]
    assert any(
        i["itemId"] == overlapping.id and i["category"] == "totals" and i["state"] == "open"
        for i in project(corrected)["issues"]
    )
    mapping = next(i for i in projection["issues"] if i.get("scope") == "table")
    accepted, _, _, _ = evaluate(
        request, operation(request, "accept", [], issueIds=[mapping["id"]])
    )
    assert set(mapping["itemIds"]) <= set(project(accepted)["readyIds"])


def test_confirmation_survives_refresh_and_an_unrelated_field_edit():
    request, _ = example_state()
    initial = synchronize(request)
    confirmation = next(i for i in initial["issues"] if i["code"] == "confirmation_required")
    assert project(request) == synchronize(request)
    assert (
        next(i for i in project(request)["issues"] if i["id"] == confirmation["id"])["code"]
        == "confirmation_required"
    )
    edited, _, _, _ = evaluate(
        request,
        operation(
            request, "edit", [confirmation["itemId"]], changes={"notes": "An unrelated note"}
        ),
    )
    assert confirmation["itemId"] not in project(edited)["readyIds"]
    edited.workflow_status, edited.review_revision = "review", 0
    accepted, _, _, _ = evaluate(
        edited, operation(edited, "accept", [], issueIds=[confirmation["id"]])
    )
    assert confirmation["itemId"] in project(accepted)["readyIds"]
    assert (
        next(i for i in project(accepted)["issues"] if i["id"] == confirmation["id"])["state"]
        == "accepted"
    )
    # A later successful whole-item AI check also satisfies the confirmation requirement.
    checked = next(i for i in request.items if i.id == confirmation["itemId"])
    checked.verification_source, checked.status, checked.confidence = "ai", "verified", 95
    assert checked.id in project(request)["readyIds"]


def test_correcting_an_invalid_zero_does_not_create_a_new_conflict_against_zero():
    request, _ = example_state()
    synchronize(request)
    invalid = next(i for i in request.items if i.attributes["Example case"] == "Invalid quantity")
    corrected, _, _, _ = evaluate(
        request, operation(request, "edit", [invalid.id], changes={"quantity": 25})
    )
    assert invalid.id in project(corrected)["readyIds"]
