import io
from typing import Any

import pytest
from httpx import ASGITransport, AsyncClient

from app.main import app

REQUEST_ID = "SD-2024-0611"


def build_xlsx(rows: list[list[Any]]) -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def build_docx(paragraphs: list[str]) -> bytes:
    from docx import Document

    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def force_llm_unavailable(monkeypatch) -> None:
    """Force extract_items_with_llm() to raise LlmUnavailable deterministically, regardless of
    what a developer's local .env happens to set (e.g. LLM_PROVIDER=gemini + a real key) -
    env vars take precedence over .env file values in pydantic-settings."""
    monkeypatch.setenv("LLM_PROVIDER", "anthropic")
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)

    from app.core.config import get_settings

    get_settings.cache_clear()


@pytest.mark.asyncio
async def test_home_returns_stats_and_recent_requests() -> None:
    response = await request("GET", "/api/home")

    assert response.status_code == 200
    body = response.json()
    assert body["organization"] == "action medeor"
    assert body["stats"][0]["label"] == "Requests Processed"
    assert body["recentRequests"][0]["id"] == REQUEST_ID


@pytest.mark.asyncio
async def test_cors_allows_localhost_alias_frontend_origin() -> None:
    response = await request(
        "OPTIONS",
        "/api/home",
        headers={
            "Origin": "http://127.0.0.1:5173",
            "Access-Control-Request-Method": "GET",
        },
    )

    assert response.status_code == 200
    assert response.headers["access-control-allow-origin"] == "http://127.0.0.1:5173"


@pytest.mark.asyncio
async def test_cors_allows_localhost_dev_server_ports() -> None:
    response = await request(
        "OPTIONS",
        "/api/home",
        headers={
            "Origin": "http://127.0.0.1:5174",
            "Access-Control-Request-Method": "GET",
        },
    )

    assert response.status_code == 200
    assert response.headers["access-control-allow-origin"] == "http://127.0.0.1:5174"


@pytest.mark.asyncio
async def test_recent_imports_returns_supported_file_types() -> None:
    response = await request("GET", "/api/imports/recent")

    assert response.status_code == 200
    body = response.json()
    assert body[0]["fileName"].endswith(".xlsx")
    assert {item["type"] for item in body} <= {"pdf", "xlsx", "xls"}


@pytest.mark.asyncio
async def test_create_import_parses_excel_and_persists_review_payload(monkeypatch) -> None:
    from app.api import routes

    async def basic_mode(*args):
        return "basic"

    # Test parser persistence independently of saved preferences or AI availability.
    monkeypatch.setattr(routes, "extraction_mode", basic_mode)
    workbook_bytes = build_xlsx(
        [
            ["Item number", "Item", "Quantity", "Unit", "desired shelf life", "Notes"],
            ["AM500-001", "Amoxicillin 500mg Capsules", 2000, "caps", "24 months", "Blister pack preferred"],
            ["PC500-001", "Paracetamol 500mg Tablets", 5000, "tabs", None, "Generic acceptable"],
            [None, "ORS Sachets (WHO formula)", None, "sachets", None, "Quantity illegible in scan"],
        ]
    )

    response = await request(
        "POST",
        "/api/imports",
        files={
            "file": (
                "Sudan_EmergencyRequest_MSF_June2024.xlsx",
                workbook_bytes,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ),
        },
    )

    assert response.status_code == 200
    body = response.json()
    assert body["requestId"].startswith("IMP-")
    assert body["source"]["fileName"] == "Sudan_EmergencyRequest_MSF_June2024.xlsx"
    assert body["counts"]["total"] == 3
    assert body["items"][0]["name"] == "Amoxicillin 500mg Capsules"
    assert body["items"][0]["quantity"] == 2000
    assert body["items"][0]["status"] == "verified"
    assert body["items"][0]["itemNumber"] == "AM500-001"
    assert body["items"][0]["shelfLife"] == "24 months"
    assert body["items"][2]["status"] == "missing"
    assert body["items"][2]["itemNumber"] == ""

    request_id = body["requestId"]

    # The freshly persisted request must be retrievable by its real requestId, not just the fixture.
    follow_up = await request("GET", f"/api/requests/{request_id}/review")
    assert follow_up.status_code == 200
    assert follow_up.json()["items"][0]["name"] == "Amoxicillin 500mg Capsules"

    # Editing itemNumber/shelfLife on a real (non-fixture) request exercises the camelCase ->
    # snake_case field mapping in repository.update_item_fields.
    item_id = body["items"][2]["id"]
    patched = await request(
        "PATCH",
        f"/api/requests/{request_id}/items/{item_id}",
        json={"quantity": 500, "itemNumber": "ORS-WHO-100", "shelfLife": "18 months"},
    )
    assert patched.status_code == 200
    patched_body = patched.json()
    assert patched_body["itemNumber"] == "ORS-WHO-100"
    assert patched_body["shelfLife"] == "18 months"
    assert patched_body["status"] == "verified"


@pytest.mark.asyncio
async def test_create_import_parses_docx_free_text(monkeypatch) -> None:
    # No LLM key available, so this exercises the naive-parse fallback end to end through the
    # real upload endpoint - not just the parser module in isolation.
    force_llm_unavailable(monkeypatch)

    docx_bytes = build_docx(
        [
            "Request from partner clinic - please supply the following:",
            "Amoxicillin 500mg caps 2000 pcs - blister pack preferred",
            "Paracetamol 500mg tablets 5000 tabs - generic acceptable",
        ]
    )

    response = await request(
        "POST",
        "/api/imports",
        files={
            "file": (
                "partner_request.docx",
                docx_bytes,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            ),
        },
    )

    assert response.status_code == 200
    body = response.json()
    assert body["requestId"].startswith("IMP-")
    assert body["counts"]["total"] == 2
    assert body["items"][0]["quantity"] == 2000
    assert body["items"][0]["unit"] == "pcs"


@pytest.mark.asyncio
async def test_create_import_parses_csv() -> None:
    csv_bytes = (
        "Item,Quantity,Unit,Notes\n"
        "Amoxicillin 500mg Capsules,2000,caps,Blister pack preferred\n"
        "Paracetamol 500mg Tablets,5000,tabs,Generic acceptable\n"
    ).encode("utf-8")

    response = await request(
        "POST",
        "/api/imports",
        files={"file": ("partner_request.csv", csv_bytes, "text/csv")},
    )

    assert response.status_code == 200
    body = response.json()
    assert body["requestId"].startswith("IMP-")
    assert body["counts"]["total"] == 2
    assert body["items"][0]["name"] == "Amoxicillin 500mg Capsules"
    assert body["items"][0]["quantity"] == 2000
    assert body["items"][0]["status"] == "verified"


@pytest.mark.asyncio
async def test_create_import_lists_skipped_columns_as_available() -> None:
    # "Supplier code"/"Unit Price" open a supplier/admin block the heuristic scopes out of the
    # initial extraction entirely - they should come back as suggestions for the post-upload
    # "Add column" control rather than just vanishing.
    csv_bytes = (
        "Item,Quantity,Unit,Supplier code,Unit Price\n"
        "Amoxicillin 500mg Capsules,2000,caps,SUP-001,1.33\n"
    ).encode("utf-8")

    response = await request(
        "POST",
        "/api/imports",
        files={"file": ("partner_request.csv", csv_bytes, "text/csv")},
    )

    assert response.status_code == 200
    body = response.json()
    assert body["availableColumns"] == ["Supplier code", "Unit Price"]
    assert body["attributeColumns"] == []
    assert body["items"][0]["attributes"] == {}


@pytest.mark.asyncio
async def test_add_custom_column_applies_via_hint_and_updates_available_columns() -> None:
    # Hint-based matching is deterministic (no LLM call needed), so this stays fast and
    # network-independent regardless of what's configured in a local .env.
    csv_bytes = (
        "Item,Quantity,Unit,Supplier code,Unit Price\n"
        "Amoxicillin 500mg Capsules,2000,caps,SUP-001,1.33\n"
    ).encode("utf-8")
    created = await request(
        "POST",
        "/api/imports",
        files={"file": ("partner_request.csv", csv_bytes, "text/csv")},
    )
    request_id = created.json()["requestId"]

    response = await request(
        "POST",
        f"/api/requests/{request_id}/custom-columns",
        json={"displayName": "Price", "hint": "Unit Price"},
    )

    assert response.status_code == 200
    body = response.json()
    assert "Price" in body["attributeColumns"]
    assert body["items"][0]["attributes"]["Price"] == "1.33"
    assert body["availableColumns"] == ["Supplier code"]


@pytest.mark.asyncio
async def test_add_custom_column_rejects_blank_name() -> None:
    csv_bytes = b"Item,Quantity,Unit\nAmoxicillin 500mg Capsules,2000,caps\n"
    created = await request(
        "POST",
        "/api/imports",
        files={"file": ("partner_request.csv", csv_bytes, "text/csv")},
    )
    request_id = created.json()["requestId"]

    response = await request(
        "POST",
        f"/api/requests/{request_id}/custom-columns",
        json={"displayName": "   ", "hint": ""},
    )

    assert response.status_code == 400


@pytest.mark.asyncio
async def test_add_custom_column_rejects_demo_request() -> None:
    response = await request(
        "POST",
        f"/api/requests/{REQUEST_ID}/custom-columns",
        json={"displayName": "Maker", "hint": ""},
    )

    assert response.status_code == 422


@pytest.mark.asyncio
async def test_update_column_label_renames_core_and_attribute_columns() -> None:
    csv_bytes = (
        "Item,Quantity,Unit,Manufacturer\n"
        "Amoxicillin 500mg Capsules,2000,caps,Reyoung Pharmaceutical\n"
    ).encode("utf-8")
    created = await request(
        "POST",
        "/api/imports",
        files={"file": ("partner_request.csv", csv_bytes, "text/csv")},
    )
    request_id = created.json()["requestId"]

    core_rename = await request(
        "PATCH",
        f"/api/requests/{request_id}/column-labels",
        json={"columnKey": "quantity", "label": "Requested Amount"},
    )
    assert core_rename.status_code == 200
    assert core_rename.json()["quantity"] == "Requested Amount"

    attribute_rename = await request(
        "PATCH",
        f"/api/requests/{request_id}/column-labels",
        json={"columnKey": "Manufacturer", "label": "Supplier Name"},
    )
    assert attribute_rename.status_code == 200
    body = attribute_rename.json()
    assert body["quantity"] == "Requested Amount"
    assert body["Manufacturer"] == "Supplier Name"

    # Persists - a fresh GET sees both renames.
    review = await request("GET", f"/api/requests/{request_id}/review")
    assert review.json()["columnLabels"] == body

    # A blank label resets that column back to its default (removes the override).
    reset = await request(
        "PATCH",
        f"/api/requests/{request_id}/column-labels",
        json={"columnKey": "quantity", "label": ""},
    )
    assert "quantity" not in reset.json()
    assert reset.json()["Manufacturer"] == "Supplier Name"


@pytest.mark.asyncio
async def test_create_import_rejects_unparsable_workbook() -> None:
    response = await request(
        "POST",
        "/api/imports",
        files={
            "file": (
                "broken.xlsx",
                b"not actually a workbook",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ),
        },
    )

    assert response.status_code == 422


@pytest.mark.asyncio
async def test_create_import_rejects_mismatched_content_type() -> None:
    response = await request(
        "POST",
        "/api/imports",
        files={
            "file": (
                "request.pdf",
                b"mock workbook bytes",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ),
        },
    )

    assert response.status_code == 400


@pytest.mark.asyncio
async def test_create_import_rejects_unsupported_extension() -> None:
    response = await request(
        "POST",
        "/api/imports",
        files={"file": ("request.txt", b"not supported", "text/plain")},
    )

    assert response.status_code == 400


@pytest.mark.asyncio
async def test_review_returns_extracted_items_and_source_references() -> None:
    response = await request("GET", f"/api/requests/{REQUEST_ID}/review")

    assert response.status_code == 200
    body = response.json()
    assert body["partner"]["partner"] == "MSF Sudan"
    assert body["sourceReferences"][0]["page"] == 3


@pytest.mark.asyncio
async def test_update_item_returns_verified_item() -> None:
    response = await request(
        "PATCH",
        f"/api/requests/{REQUEST_ID}/items/3",
        json={"quantity": 500, "notes": "Quantity confirmed by phone"},
    )

    assert response.status_code == 200
    body = response.json()
    assert body["quantity"] == 500
    assert body["status"] == "verified"


@pytest.mark.asyncio
async def test_verify_item_returns_verified_item() -> None:
    response = await request("POST", f"/api/requests/{REQUEST_ID}/items/5/verify")

    assert response.status_code == 200
    assert response.json()["status"] == "verified"


@pytest.mark.asyncio
async def test_update_partner_requires_separate_confirmation() -> None:
    response = await request(
        "PATCH",
        f"/api/requests/{REQUEST_ID}/partner",
        json={
            "partner": "MSF Sudan",
            "region": "East Sudan (Kassala)",
            "requestId": REQUEST_ID,
            "contact": "Dr. Amira Hassan",
        },
    )

    assert response.status_code == 200
    assert response.json()["confirmed"] is False


@pytest.mark.asyncio
async def test_start_matching_returns_candidates_and_defaults() -> None:
    response = await request("POST", f"/api/requests/{REQUEST_ID}/matching")

    assert response.status_code == 200
    body = response.json()
    assert body["requestedItems"][0]["name"] == "Amoxicillin 500mg Capsules"
    assert body["matches"]["1"][0]["id"] == "erp-001"
    assert body["selectedMatches"]["1"] == "erp-001"


@pytest.mark.asyncio
async def test_update_matching_returns_selected_match() -> None:
    response = await request(
        "PATCH",
        f"/api/requests/{REQUEST_ID}/matching/1",
        json={"matchId": "erp-002"},
    )

    assert response.status_code == 200
    assert response.json() == {"itemId": 1, "matchId": "erp-002"}


@pytest.mark.asyncio
async def test_summary_returns_items_and_metrics() -> None:
    response = await request("GET", f"/api/requests/{REQUEST_ID}/summary")

    assert response.status_code == 200
    body = response.json()
    assert body["items"][0]["sku"] == "AM500-CAP-1000"
    assert body["metrics"]["totalLineItems"] == 8


@pytest.mark.asyncio
async def test_create_offer_returns_mock_document_details() -> None:
    response = await request("POST", f"/api/requests/{REQUEST_ID}/offer")

    assert response.status_code == 200
    body = response.json()
    assert body["fileName"] == f"Offer-{REQUEST_ID}.pdf"
    assert body["lineItems"] == 8


@pytest.mark.asyncio
async def test_trends_returns_chart_data() -> None:
    response = await request("GET", "/api/trends")

    assert response.status_code == 200
    body = response.json()
    assert body["demandTrend"][0]["month"] == "Jan"
    assert body["categoryDemand"][0]["risk"] == "high"
    assert body["topItems"][0]["name"] == "Paracetamol 500mg Tablets"


@pytest.mark.asyncio
async def test_unknown_request_returns_404() -> None:
    response = await request("GET", "/api/requests/UNKNOWN/review")

    assert response.status_code == 404


async def request(method: str, url: str, **kwargs: Any) -> Any:
    async with AsyncClient(
        transport=ASGITransport(app=app),
        base_url="http://testserver",
    ) as client:
        return await client.request(method, url, **kwargs)


@pytest.mark.asyncio
async def test_delete_request_requires_a_saved_request_and_rejects_active_matching(monkeypatch) -> None:
    from app.api import routes
    from app.db import repository
    from app.db.session import get_session

    async def fake_session():
        yield object()

    async def fake_delete(_session, request_id: str) -> bool:
        if request_id == "ACTIVE":
            raise repository.RequestCurrentlyMatching()
        return request_id == "SAVED"

    app.dependency_overrides[get_session] = fake_session
    monkeypatch.setattr(routes.repository, "delete_request", fake_delete)
    try:
        assert (await request("DELETE", "/api/requests/SAVED")).status_code == 204
        assert (await request("DELETE", "/api/requests/MISSING")).status_code == 404
        active = await request("DELETE", "/api/requests/ACTIVE")
        assert active.status_code == 409
        assert "matching finishes" in active.json()["detail"]
    finally:
        app.dependency_overrides.pop(get_session, None)


def build_two_sheet_xlsx() -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    first = workbook.active
    first.title = "Request"
    for row in [["Product", "Quantity", "Unit"], ["Sterile gloves", 200, "pairs"]]:
        first.append(row)
    other = workbook.create_sheet("Second list")
    for row in [["Product", "Quantity", "Unit"], ["Syringe 5 ml", 30, "pcs"], ["Gauze", 10, "rolls"]]:
        other.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.mark.asyncio
async def test_worksheet_can_be_switched_during_review_keeping_manual_rows() -> None:
    created = await request(
        "POST", "/api/imports", files={"file": ("two.xlsx", build_two_sheet_xlsx(), XLSX)}
    )
    body = created.json()
    request_id = body["requestId"]
    assert body["selectedSheet"] == "Request"
    assert body["sheets"] == [{"name": "Request", "rows": 2}, {"name": "Second list", "rows": 3}]
    assert [item["name"] for item in body["items"]] == ["Sterile gloves"]

    manual = await request(
        "POST", f"/api/requests/{request_id}/items",
        json={"name": "Hand-added mask", "quantity": 5, "unit": "pcs", "domain": "equipment"},
    )
    assert manual.status_code == 201

    switched = await request(
        "POST", f"/api/requests/{request_id}/worksheet", json={"sheet": "Second list"}
    )
    assert switched.status_code == 200
    body = switched.json()
    assert body["selectedSheet"] == "Second list"
    assert [item["name"] for item in body["items"]] == ["Syringe 5 ml", "Gauze", "Hand-added mask"]

    # Later re-parses (here: adding a column) read the selected sheet, not the first one.
    with_column = await request(
        "POST", f"/api/requests/{request_id}/custom-columns",
        json={"displayName": "Unit copy", "hint": "Unit"},
    )
    assert [item["name"] for item in with_column.json()["items"]][:2] == ["Syringe 5 ml", "Gauze"]


@pytest.mark.asyncio
async def test_unknown_worksheet_is_rejected() -> None:
    created = await request(
        "POST", "/api/imports", files={"file": ("two.xlsx", build_two_sheet_xlsx(), XLSX)}
    )
    response = await request(
        "POST", f"/api/requests/{created.json()['requestId']}/worksheet", json={"sheet": "Nope"}
    )
    assert response.status_code == 400


@pytest.mark.asyncio
async def test_upload_extracts_only_confirmed_worksheets(monkeypatch):
    from app.api import routes
    from tests.test_excel_first_sheet import workbook_bytes

    async def basic_mode(*args):
        return "basic"

    monkeypatch.setattr(routes, "extraction_mode", basic_mode)
    content = workbook_bytes(
        [["Product", "Quantity", "Unit"], ["First device", 2, "pcs"]],
        [["Product", "Quantity", "Unit"], ["Other device", 9, "pcs"]],
    )
    for selection, expected in [
        (["Other request"], ["Other device"]),
        (["First sheet", "Other request"], ["First device", "Other device"]),
    ]:
        response = await request("POST", "/api/imports",
                                 files={"file": ("two.xlsx", content)},
                                 data={"selected_sheets": selection, "progress_token": "worksheet-merge-test"})
        assert response.status_code == 200
        body = response.json()
        assert [item["name"] for item in body["items"]] == expected
        assert body["selectedSheets"] == selection
        saved = await request("GET", f'/api/requests/{body["requestId"]}/review')
        assert saved.json()["selectedSheets"] == selection


@pytest.mark.asyncio
async def test_upload_reports_real_stages_and_clears_them(monkeypatch) -> None:
    from app.api import import_progress, routes

    stages = []
    real_report = import_progress.report

    async def recording(token, stage, **kwargs):
        stages.append((stage, kwargs.get("items")))
        await real_report(token, stage, **kwargs)
        # While the upload runs, the poll endpoint sees the stage just written.
        polled = await request("GET", f"/api/imports/progress/{token}")
        assert polled.json()["stage"] == stage

    monkeypatch.setattr(routes.import_progress, "report", recording)
    csv_bytes = b"Product,Quantity,Unit\nGauze,5,rolls\nGloves,10,pairs\n"
    token = "test-token-0001"
    response = await request(
        "POST", "/api/imports",
        files={"file": ("request.csv", csv_bytes, "text/csv")},
        data={"progress_token": token},
    )
    assert response.status_code == 200
    assert stages[0] == ("reading", None)
    assert stages[-1] == ("saving", 2)
    # Finished uploads leave nothing behind to poll.
    assert (await request("GET", f"/api/imports/progress/{token}")).status_code == 404


@pytest.mark.asyncio
async def test_progress_for_unknown_or_malformed_token_is_404() -> None:
    assert (await request("GET", "/api/imports/progress/never-started-123")).status_code == 404
    assert (await request("GET", "/api/imports/progress/bad%20token")).status_code == 404
