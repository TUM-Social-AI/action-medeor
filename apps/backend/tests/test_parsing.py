import io

from app.parsing.docx_parser import parse_docx
from app.parsing.excel_parser import parse_excel
from app.parsing.keywords import match_column_role
from app.parsing.pdf_parser import parse_pdf
from app.parsing.table_parser import is_table_well_structured, parse_table_rows
from app.parsing.text_heuristics import classify_item, extract_quantity_and_unit, parse_number


def force_llm_unavailable(monkeypatch) -> None:
    """Force extract_items_with_llm() to raise LlmUnavailable deterministically, regardless of
    what a developer's local .env happens to set (e.g. LLM_PROVIDER=gemini + a real key) -
    env vars take precedence over .env file values in pydantic-settings, so this is reliable
    even when the process was started with a fully configured .env on disk."""
    monkeypatch.setenv("LLM_PROVIDER", "anthropic")
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)

    from app.core.config import get_settings

    get_settings.cache_clear()


def build_xlsx(rows: list[list[object]]) -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def build_minimal_pdf(lines: list[str]) -> bytes:
    """Hand-rolled single-page PDF with a Helvetica text stream - no PDF-writing dependency
    is installed, and this is enough to exercise pdfplumber's real text extraction."""

    def escape(text: str) -> str:
        return text.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")

    text_ops = "BT /F1 12 Tf 14 TL 50 750 Td\n"
    for line in lines:
        text_ops += f"({escape(line)}) Tj T*\n"
    text_ops += "ET"
    stream = text_ops.encode("latin-1")

    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /Resources << /Font << /F1 4 0 R >> >> "
        b"/MediaBox [0 0 612 792] /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        f"<< /Length {len(stream)} >>\nstream\n".encode("latin-1") + stream + b"\nendstream",
    ]

    out = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{index} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref_offset = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for offset in offsets[1:]:
        out += f"{offset:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF".encode()
    return bytes(out)


def build_docx(paragraphs: list[str], table_rows: list[list[str]] | None = None) -> bytes:
    from docx import Document

    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)

    if table_rows:
        table = document.add_table(rows=0, cols=len(table_rows[0]))
        for row_values in table_rows:
            row = table.add_row()
            for cell, value in zip(row.cells, row_values, strict=True):
                cell.text = value

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_match_column_role_recognizes_multiple_languages() -> None:
    assert match_column_role("Item Description") == "name"
    assert match_column_role("Bezeichnung") == "name"
    assert match_column_role("Désignation") == "name"
    assert match_column_role("الصنف") == "name"
    assert match_column_role("Qty") == "quantity"
    assert match_column_role("Menge") == "quantity"
    assert match_column_role("Unit") == "unit"
    assert match_column_role("Some Random Header") is None


def test_parse_number_handles_various_shapes() -> None:
    assert parse_number(2000) == 2000
    assert parse_number(2000.0) == 2000
    assert parse_number("2,000") == 2000
    assert parse_number("2000 pcs") == 2000
    assert parse_number(None) is None
    assert parse_number("") is None


def test_extract_quantity_and_unit_from_free_text() -> None:
    quantity, unit = extract_quantity_and_unit("Amoxicillin 500mg caps [2000 pcs]")
    assert quantity == 2000
    assert unit == "pcs"


def test_classify_item_statuses() -> None:
    assert classify_item("", None, "")[0] == "missing"
    assert classify_item("Item", None, "")[0] == "missing"
    assert classify_item("Item", 10, "")[0] == "needs_review"
    assert classify_item("Item", 10, "caps")[0] == "verified"


def test_parse_table_rows_detects_header_and_maps_columns() -> None:
    rows = [
        ["Item", "Quantity", "Unit", "Notes"],
        ["Amoxicillin 500mg Capsules", 2000, "caps", "Blister pack preferred"],
        ["ORS Sachets", None, "sachets", "Illegible"],
    ]

    result = parse_table_rows(rows)

    assert result.rows_detected == 2
    assert result.items[0].name == "Amoxicillin 500mg Capsules"
    assert result.items[0].quantity == 2000
    assert result.items[0].status == "verified"
    assert result.items[1].quantity is None
    assert result.items[1].status == "missing"


def test_parse_table_rows_does_not_confuse_item_number_with_name() -> None:
    # Real RFQ-tracker workbooks have an "Item number" (SKU) column before the actual
    # description column ("Product") - "Item number" must not steal the name role via a bare
    # "item" keyword match, and its value should land in the dedicated item_number field.
    rows = [
        ["Item number", "Product", "Quantity enquiry", "Packaging enquiry"],
        [424109001, "Cholera Tests, stool, rapid", 300, None],
        [None, "Dengue Test", 200, None],
    ]

    result = parse_table_rows(rows)

    assert result.rows_detected == 2
    assert result.items[0].name == "Cholera Tests, stool, rapid"
    assert result.items[0].item_number == "424109001"
    assert result.items[1].name == "Dengue Test"
    assert result.items[1].item_number == ""


def test_parse_table_rows_captures_translation_and_shelf_life() -> None:
    rows = [
        ["Product", "Translation", "desired shelf life", "Quantity enquiry"],
        ["water-based lubricant, sterile", "Wasserbasierendes Gleitmittel, steril", "16 months", 272],
    ]

    result = parse_table_rows(rows)

    assert result.rows_detected == 1
    item = result.items[0]
    assert item.shelf_life == "16 months"
    assert "Wasserbasierendes Gleitmittel" in item.notes


def test_extract_request_priority_hint_reads_special_info_field() -> None:
    from app.parsing.table_parser import extract_request_priority_hint

    rows = [
        [None, "Besondere Informationen:", "DRINGEND - komplette Lieferung notwendig"],
        ["Product", "Quantity enquiry"],
    ]

    assert extract_request_priority_hint(rows) == "critical"


def test_extract_request_priority_hint_returns_none_when_empty() -> None:
    from app.parsing.table_parser import extract_request_priority_hint

    rows = [
        [None, "Besondere Informationen:", None],
        ["Product", "Quantity enquiry"],
    ]

    assert extract_request_priority_hint(rows) is None


def test_parse_table_rows_ignores_supplier_quote_block() -> None:
    # Everything from "Supplier I" onward is action medeor's own quote tracking, not part of
    # what the partner requested - must not bleed into quantity/unit/notes or the excerpt.
    rows = [
        ["Product", "Quantity enquiry", "Packaging enquiry", "Supplier I", "Item offered", "Quantity on offer", "Packaging"],
        ["Cholera Tests, stool, rapid", 300, None, "Centramed", "Rapid Test Kit, 20/box", 15, 20],
    ]

    result = parse_table_rows(rows)

    assert result.rows_detected == 1
    item = result.items[0]
    assert item.quantity == 300
    assert item.unit == ""
    assert "Centramed" not in item.excerpt
    assert "Rapid Test Kit" not in item.excerpt
    assert any("supplier" in warning.lower() for warning in result.warnings)


def test_parse_table_rows_skips_blank_names() -> None:
    rows = [
        ["Item", "Quantity", "Unit"],
        ["", 100, "caps"],
        ["Paracetamol 500mg Tablets", 5000, "tabs"],
    ]

    result = parse_table_rows(rows)

    assert len(result.items) == 1
    assert result.items[0].name == "Paracetamol 500mg Tablets"


def test_is_table_well_structured_true_for_consistent_table() -> None:
    rows = [
        ["Item", "Quantity", "Unit"],
        ["Amoxicillin 500mg Capsules", 2000, "caps"],
        ["Paracetamol 500mg Tablets", 5000, "tabs"],
        ["ORS Sachets", 500, "sachets"],
    ]
    assert is_table_well_structured(rows) is True


def test_is_table_well_structured_false_without_header() -> None:
    rows = [
        ["Please find our request below."],
        ["We urgently need medical supplies for the region."],
    ]
    assert is_table_well_structured(rows) is False


def test_parse_excel_end_to_end() -> None:
    content = build_xlsx(
        [
            ["Item", "Quantity", "Unit", "Notes"],
            ["Wound Dressing 10x10cm", 1000, "pcs", "Sterile"],
            ["Morphine HCl 10mg/ml Injection", 50, "vials", "Controlled substance"],
        ]
    )

    document = parse_excel(content, "request.xlsx")

    assert document.rows_detected == 2
    assert {item.name for item in document.items} == {
        "Wound Dressing 10x10cm",
        "Morphine HCl 10mg/ml Injection",
    }


def test_parse_excel_falls_back_to_free_text_quantity_unit() -> None:
    # Some files put the whole "name + qty + unit" into a single description column.
    content = build_xlsx(
        [
            ["Description"],
            ["Amoxicillin 500mg caps 2000 pcs"],
        ]
    )

    document = parse_excel(content, "request.xlsx")

    assert document.rows_detected == 1
    assert document.items[0].quantity == 2000
    assert document.items[0].unit == "pcs"


def test_parse_pdf_free_text_falls_back_to_naive_parsing_without_llm_key(monkeypatch) -> None:
    # No LLM key configured, so the free-text page (no detectable table) should degrade to the
    # naive per-line regex parser, not raise.
    force_llm_unavailable(monkeypatch)

    content = build_minimal_pdf(
        [
            "Amoxicillin 500mg caps 2000 pcs - blister pack preferred",
            "Paracetamol 500mg tablets 5000 tabs - generic acceptable",
        ]
    )

    document = parse_pdf(content)

    assert document.used_llm_fallback is False
    assert any("LLM fallback unavailable" in warning for warning in document.warnings)
    assert document.rows_detected == 2
    assert document.items[0].quantity == 2000
    assert document.items[0].unit == "pcs"
    assert document.items[0].confidence == 40


def test_parse_docx_falls_back_to_naive_parsing_without_llm_key(monkeypatch) -> None:
    # .docx always goes straight to the free-text path (no table-detection step), so this
    # exercises the same LLM -> naive fallback as PDF free text, just via a different reader.
    force_llm_unavailable(monkeypatch)

    content = build_docx(
        [
            "Partner request - urgent medical resupply",
            "Amoxicillin 500mg caps 2000 pcs - blister pack preferred",
            "Paracetamol 500mg tablets 5000 tabs - generic acceptable",
        ]
    )

    document = parse_docx(content)

    assert document.used_llm_fallback is False
    assert any("LLM fallback unavailable" in warning for warning in document.warnings)
    assert document.rows_detected == 2
    assert document.items[0].quantity == 2000
    assert document.items[0].unit == "pcs"
    assert document.items[0].confidence == 40


def test_parse_docx_includes_table_text(monkeypatch) -> None:
    force_llm_unavailable(monkeypatch)

    content = build_docx(
        ["Supplies needed for the clinic:"],
        table_rows=[["Wound Dressing 10x10cm 1000 pcs sterile"]],
    )

    document = parse_docx(content)

    assert document.rows_detected == 1
    assert document.items[0].quantity == 1000


def test_parse_docx_empty_document_warns() -> None:
    content = build_docx([])

    document = parse_docx(content)

    assert document.rows_detected == 0
    assert any("no readable text" in warning for warning in document.warnings)


PUI_HEADER = [
    "No",
    "Code PUI",
    "Designation",
    "Unit Qty / Qté unitaire (ex: 1 comprimé, 1 ampoule etc)",
    "Supplier code / Code fournisseur",
    "Pack Size / Packaging",
    "Unit Price / Prix à l'unité",
    "Manufacturer Name / Nom du Fabricant",
    "Vendors comments",
    "Article priority / Priorité d'article",
    "Quality assurance requirements for validation",
    "Additional documents required",
    "Comments / Commentaires",
]


def test_parse_table_rows_handles_supplier_block_in_the_middle() -> None:
    # Premiere Urgence's form sandwiches the supplier's section between the request columns and
    # the requester's own fields, unlike the Anfrage trackers which append it at the end. The
    # requester fields after the block must survive.
    rows = [
        PUI_HEADER,
        [1, "DORAACSA7TG", "Acid ACETYLSALICYLIC, 75mg, tab.", 25330, "209126003", 100, 1.33,
         "Reyoung Pharmaceutical Co., Ltd.", "some vendor note", "Standard",
         "Source mandatory (manufacturer+country)", None, "Must be WHO prequalified"],
    ]

    result = parse_table_rows(rows)

    assert result.rows_detected == 1
    item = result.items[0]
    # "Designation" must win the name role even though it lacks the accent our keyword carries.
    assert item.name == "Acid ACETYLSALICYLIC, 75mg, tab."
    assert item.quantity == 25330
    assert item.item_number == "DORAACSA7TG"
    # Nothing from the supplier's section leaks into the item or its excerpt.
    assert "Reyoung" not in item.excerpt
    assert "vendor" not in item.excerpt.lower()
    # ...while the requester's own columns after that block are kept. "Article priority" maps
    # onto our scale (this partner's "Standard" tier -> medium) rather than staying a raw
    # attribute - see test_procurement_priority_is_mapped_not_kept_raw for the mapping itself.
    assert item.priority == "medium"
    assert "Source mandatory" in item.attributes["Quality assurance requirements for validation"]
    assert item.notes == "Must be WHO prequalified"


def test_parse_table_rows_drops_ordinal_and_all_empty_columns() -> None:
    rows = [
        PUI_HEADER,
        [1, "A1", "Item one", 10, None, None, None, None, None, "Standard", "QA note", None, None],
        [2, "A2", "Item two", 20, None, None, None, None, None, "Standard", "QA note", None, None],
    ]

    result = parse_table_rows(rows)

    # "No" is a bare row ordinal (redundant with the table's own numbering), "Additional
    # documents required" is empty on every row, and "Article priority" is mapped onto the
    # priority field rather than kept as a raw column - none of the three earns an attribute.
    assert result.attribute_columns == ["Quality assurance requirements for validation"]


def test_unrecognized_priority_wording_is_kept_raw_instead_of_forced() -> None:
    rows = [
        PUI_HEADER,
        [1, "A1", "Item one", 10, None, None, None, None, None, "Confidential - do not disclose",
         "QA note", None, None],
    ]

    result = parse_table_rows(rows)
    item = result.items[0]

    # This wording has no equivalent on our critical/high/medium/low scale (nor in the
    # procurement-tier vocabulary), so the priority field stays at its default and the partner's
    # own wording is preserved verbatim instead of being force-fitted or dropped.
    assert item.priority == "medium"
    assert item.attributes["Article priority / Priorité d'article"] == "Confidential - do not disclose"


def test_accent_folding_matches_headers_written_without_accents() -> None:
    assert match_column_role("Designation") == "name"
    assert match_column_role("Désignation") == "name"
    assert match_column_role("Qte") == "quantity"
    assert match_column_role("Qté") == "quantity"
    assert match_column_role("Unite") == "unit"


def test_detect_priority_from_column_value_maps_procurement_tiers() -> None:
    from app.parsing.text_heuristics import detect_priority_from_column_value

    assert detect_priority_from_column_value("Prioritaire-Priority") == "high"
    assert detect_priority_from_column_value("Standard") == "medium"
    assert detect_priority_from_column_value("Optionnel-Optional") == "low"
    assert detect_priority_from_column_value("Alternative Standard") == "low"
    # A column that already spells out our own vocabulary still works.
    assert detect_priority_from_column_value("Critical") == "critical"
    assert detect_priority_from_column_value("something unrecognized") is None


def test_procurement_priority_is_mapped_not_kept_raw() -> None:
    rows = [
        ["Product", "Quantity enquiry", "Article priority"],
        ["Item one", 10, "Prioritaire-Priority"],
    ]

    result = parse_table_rows(rows)
    item = result.items[0]

    assert item.priority == "high"
    # Once mapped, the raw value isn't duplicated into attributes.
    assert "Article priority" not in item.attributes


def test_parse_table_rows_computes_total_from_packs_times_units_per_pack() -> None:
    rows = [
        ["Product", "Packs Requested", "Units Per Pack"],
        ["Item one", 15, 100],
    ]
    # Header keywords alone won't map "Packs Requested"/"Units Per Pack" (that mapping normally
    # comes from LLM classification - see test_llm_table_classifier.py); this test exercises the
    # deterministic multiplication logic in parse_table_rows() directly via an injected layout.
    from app.parsing.table_parser import HeaderLayout

    layout = HeaderLayout(
        row_index=0,
        roles={0: "name", 1: "quantity_packs", 2: "units_per_pack"},
        labels={0: "Product", 1: "Packs Requested", 2: "Units Per Pack"},
    )

    result = parse_table_rows(rows, layout=layout)
    item = result.items[0]

    assert item.quantity == 1500
    assert item.attributes["Packs requested"] == "15"
    assert item.attributes["Units per pack"] == "100"


def test_parse_csv_comma_delimited() -> None:
    from app.parsing.csv_parser import parse_csv

    content = (
        "Item,Quantity,Unit,Notes\n"
        "Amoxicillin 500mg Capsules,2000,caps,Blister pack preferred\n"
        "ORS Sachets,,sachets,Illegible\n"
    ).encode("utf-8")

    document = parse_csv(content)

    assert document.rows_detected == 2
    assert document.items[0].name == "Amoxicillin 500mg Capsules"
    assert document.items[0].quantity == 2000
    assert document.items[0].status == "verified"
    assert document.items[1].quantity is None
    assert document.items[1].status == "missing"


def test_parse_csv_semicolon_delimited_with_accents() -> None:
    from app.parsing.csv_parser import parse_csv

    # European export style: semicolon delimiter, cp1252 encoding (comma is the decimal
    # separator in these locales, so semicolon is the common CSV delimiter instead).
    content = (
        "Désignation;Quantité;Unité\n"
        "Amoxicilline 500mg comprimés;2000;comprimés\n"
    ).encode("cp1252")

    document = parse_csv(content)

    assert document.rows_detected == 1
    assert document.items[0].name == "Amoxicilline 500mg comprimés"
    assert document.items[0].quantity == 2000


def test_parse_csv_empty_file_warns() -> None:
    from app.parsing.csv_parser import parse_csv

    document = parse_csv(b"")

    assert document.rows_detected == 0
    assert any("no rows" in warning for warning in document.warnings)
