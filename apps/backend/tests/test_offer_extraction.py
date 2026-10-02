"""Offline quotation tests: real document readers, controlled LLM responses, no API/DB."""

import importlib.util
import io
import json
from pathlib import Path

import pytest
from openpyxl import Workbook
from pydantic import ValidationError
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.pdfgen import canvas
from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, TableStyle

from app.offers.extraction import (
    Evidence,
    ExtractedOffer,
    _LlmResult,
    extract_offers,
    read_offer_document,
)
from app.parsing.llm_client import LlmUnavailable


def workbook_bytes(sheets: dict[str, list[list[object]]]) -> bytes:
    book = Workbook()
    book.remove(book.active)
    for name, rows in sheets.items():
        sheet = book.create_sheet(name)
        for row in rows:
            sheet.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    book.close()
    return buffer.getvalue()


def offer(**overrides) -> ExtractedOffer:
    values = {
        "source_id": "Offers!A2",
        "alternative_index": 1,
        "supplier": "Supplier A",
        "item_description": "Needle 19G",
        "offer_reference": None,
        "offer_date": None,
        "date_kind": None,
        "valid_until": None,
        "validity_text": None,
        "price_text": "42,00 EUR / 100 Stück",
        "price_amount": "42.00",
        "currency": "EUR",
        "price_basis": "100 Stück",
        "price_conflict": False,
        "warnings": [],
        "evidence": [
            Evidence(field="supplier", location="Offers!A2", excerpt="Supplier A"),
            Evidence(field="item_description", location="Offers!B2", excerpt="Needle 19G"),
            Evidence(field="price", location="Offers!C2", excerpt="42,00 EUR / 100 Stück"),
        ],
    }
    return ExtractedOffer.model_validate(values | overrides)


def test_reader_keeps_all_sheets_coordinates_headers_formats_and_formula_cache() -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "Offers"
    sheet.append(["Supplier I", "Item offered", "Price"])
    sheet.append(["Supplier A", "Needle 19G", 42])
    sheet["C2"].number_format = '#,##0.00 "EUR"'
    sheet["D2"] = "=C2*2"
    sheet.merge_cells("A2:A3")
    book.create_sheet("More offers").append(["Supplier B", "Syringe 5 ml"])
    buffer = io.BytesIO()
    book.save(buffer)
    book.close()
    chunks, warnings = read_offer_document(buffer.getvalue(), "offers.xlsx")
    text = "\n".join(chunk.text for chunk in chunks)
    assert {c.label for c in chunks} == {"Offers", "More offers"}
    assert "A2:A3" in text and '"C2"' in text
    assert '"source_ids": ["Offers!A2"]' in text
    assert "EUR" in text and "Supplier B" in text
    assert "=C2*2" in text and "No cached formula value" in text
    assert warnings == []


def test_complete_rows_are_split_with_repeated_headers_and_no_loss() -> None:
    content = workbook_bytes(
        {
            "Offers": [["Supplier I", "Item offered"]]
            + [[f"Supplier {n}", f"Needle {n}"] for n in range(80)]
        }
    )
    chunks, _ = read_offer_document(content, "offers.xlsx", max_chars=1000)
    assert len(chunks) > 1
    rows = []
    for chunk in chunks:
        assert len(chunk.text) <= 1000
        lines = chunk.text.splitlines()
        assert "Supplier I" in lines[0]
        rows.extend(json.loads(line)["row"] for line in lines[1:])
    assert rows == list(range(2, 82))


def test_oversized_single_row_fails_without_silent_truncation_or_model_call() -> None:
    def unexpected(*args):
        pytest.fail("No LLM request should happen for incompletely readable input")

    content = workbook_bytes({"Offers": [["Supplier I", "Item offered"], ["A", "x" * 4000]]})
    result = extract_offers(content, "offers.xlsx", llm=unexpected, max_chars=500)
    assert result.failures and "complete row exceeds" in result.failures[0]
    assert result.chunks_attempted == 0


def test_distinct_supplier_blocks_and_alternatives_remain_distinct_without_price_math() -> None:
    content = workbook_bytes(
        {
            "Offers": [
                ["Supplier I", "Item offered", "Price", "Supplier II", "Item offered", "Price"],
                [
                    "Supplier A",
                    "Needle 19G",
                    "42,00 EUR / 100 Stück",
                    "Supplier B",
                    "Needle 21G",
                    "10 EUR / 4 Stück",
                ],
            ]
        }
    )
    a = offer()
    b = offer(
        source_id="Offers!D2",
        supplier="Supplier B",
        item_description="Needle 21G",
        price_amount="10",
        price_basis="4 Stück",
        price_text="10 EUR / 4 Stück",
    )
    alt = offer(
        source_id="Offers!D2",
        alternative_index=2,
        supplier="Supplier B",
        item_description="Needle 23G",
        price_amount=None,
        price_text=None,
        price_basis=None,
        currency=None,
    )
    result = extract_offers(
        content, "offers.xlsx", llm=lambda *_: _LlmResult(offers=[a, b, alt], warnings=[])
    )
    assert len(result.offers) == 3
    assert [o.price_amount for o in result.offers] == ["42.00", "10", None]
    assert [o.price_basis for o in result.offers] == ["100 Stück", "4 Stück", None]


def test_conflicting_prices_cannot_select_an_amount_and_both_raw_values_survive() -> None:
    conflicted = offer(
        price_amount="27.258",
        price_conflict=True,
        evidence=[
            Evidence(field="price", location="Offers!C2", excerpt="27,258€ / 42 Stück"),
            Evidence(field="price", location="Offers!D2", excerpt="27.257"),
        ],
    )
    assert conflicted.price_amount is None
    assert len(conflicted.evidence) == 2
    assert conflicted.warnings == ["Conflicting quoted prices require review"]


@pytest.mark.parametrize("value", ["NaN", "Infinity", "-1", "42,00", "free"])
def test_invalid_quote_amounts_rejected(value: str) -> None:
    with pytest.raises(ValidationError):
        offer(price_amount=value)


def test_relative_offer_validity_is_preserved_without_date_arithmetic_or_fallback() -> None:
    content = workbook_bytes(
        {"Offers": [["Supplier I", "Item offered"], ["Supplier A", "Needle 19G"]]}
    )
    original = offer(
        offer_date="2026-09-10",
        date_kind="issued",
        valid_until=None,
        validity_text="Offer valid for 30 days from issue",
    )
    extracted = extract_offers(
        content, "offers.xlsx", llm=lambda *_: _LlmResult(offers=[original], warnings=[])
    )
    assert extracted.offers[0].offer_date.isoformat() == "2026-09-10"
    assert extracted.offers[0].valid_until is None
    original = offer()
    extracted = extract_offers(
        content, "offers.xlsx", llm=lambda *_: _LlmResult(offers=[original], warnings=[])
    )
    assert extracted.offers[0].offer_date is None
    assert extracted.offers[0].valid_until is None


def test_model_failure_keeps_partial_results_and_reports_failed_chunk() -> None:
    content = workbook_bytes(
        {
            "Offers": [["Supplier I", "Item offered"]]
            + [["Supplier A", "Needle 19G"] for _ in range(25)]
        }
    )
    calls = 0

    def controlled(*args):
        nonlocal calls
        calls += 1
        if calls == 2:
            raise LlmUnavailable("test outage")
        return _LlmResult(offers=[offer()], warnings=[])

    result = extract_offers(content, "offers.xlsx", llm=controlled)
    assert len(result.offers) == 1
    assert result.chunks_attempted == 2 and result.chunks_succeeded == 1
    assert "test outage" in result.failures[0]


def test_corrupt_and_unsupported_files_are_explicit_failures() -> None:
    for filename in ("broken.pdf", "broken.xlsx", "conversation.eml"):
        result = extract_offers(b"not a document", filename)
        assert result.failures and result.offers == []


def test_blank_pdf_reports_unsupported_without_calling_model() -> None:
    buffer = io.BytesIO()
    writer = canvas.Canvas(buffer)
    writer.showPage()
    writer.save()
    result = extract_offers(buffer.getvalue(), "empty.pdf", llm=lambda *_: pytest.fail("No text"))
    assert result.chunks_attempted == 0
    assert "scanned/empty page unsupported" in result.warnings[0]


def test_blank_workbook_does_not_call_model() -> None:
    result = extract_offers(
        workbook_bytes({"Blank": []}),
        "empty.xlsx",
        llm=lambda *_: pytest.fail("No populated cells"),
    )
    assert result.chunks_attempted == 0
    assert result.warnings == ["Workbook contained no populated cells"]


def test_fabricated_source_location_and_excerpt_are_flagged() -> None:
    content = workbook_bytes(
        {"Offers": [["Supplier I", "Item offered"], ["Supplier A", "Needle 19G"]]}
    )
    original = offer(
        evidence=[
            Evidence(field="supplier", location="Offers!A2", excerpt="Invented supplier"),
            Evidence(field="item_description", location="Offers!Z999", excerpt="Needle 19G"),
        ]
    )
    result = extract_offers(
        content, "offers.xlsx", llm=lambda *_: _LlmResult(offers=[original], warnings=[])
    )
    warnings = result.offers[0].warnings
    assert "Unverified source evidence: supplier at Offers!A2" in warnings
    assert "Unverified source evidence: item_description at Offers!Z999" in warnings
    assert "Missing source evidence for: price" in warnings


def test_placeholder_sheet_name_in_identity_is_flagged() -> None:
    content = workbook_bytes(
        {"Offers": [["Supplier I", "Item offered"], ["Supplier A", "Needle 19G"]]}
    )
    original = offer(source_id="SHEET!A2")
    result = extract_offers(
        content, "offers.xlsx", llm=lambda *_: _LlmResult(offers=[original], warnings=[])
    )
    assert "Unverified source identity: SHEET!A2" in result.offers[0].warnings


def test_pdf_reader_keeps_header_footer_tables_and_continuation_context() -> None:
    buffer = io.BytesIO()
    writer = canvas.Canvas(buffer)
    writer.drawString(40, 750, "Supplier A, issued 2026-09-10, valid until 2026-11-30")
    writer.drawString(40, 700, "1 Needle 19G, EUR 42 / 100 pieces")
    writer.drawString(40, 40, "Footer quote Q123")
    writer.showPage()
    writer.drawString(40, 750, "Continued: 2 Syringe 5 ml, EUR 9 / box")
    writer.save()
    chunks, _ = read_offer_document(buffer.getvalue(), "offers.pdf")
    text = "\n".join(c.text for c in chunks)
    assert "Footer quote Q123" in text and "Continued" in text
    assert "2026-11-30" in text and '"tables"' in text


def test_wrapped_pdf_table_description_is_valid_source_evidence() -> None:
    buffer = io.BytesIO()
    description = "Sterile needle 19G with a very long descriptive product name"
    table = Table(
        [["Item", "Qty"], [Paragraph(description, getSampleStyleSheet()["Normal"]), "500"]],
        colWidths=[130, 70],
    )
    table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 1, "black")]))
    SimpleDocTemplate(buffer).build([table])
    original = offer(
        source_id="page:1:line:1",
        supplier=None,
        item_description=description,
        price_text=None,
        price_amount=None,
        currency=None,
        price_basis=None,
        evidence=[Evidence(field="item_description", location="page:1", excerpt=description)],
    )
    result = extract_offers(
        buffer.getvalue(), "wrapped.pdf", llm=lambda *_: _LlmResult(offers=[original], warnings=[])
    )
    assert not any("source evidence" in warning for warning in result.offers[0].warnings)


def test_scoring_separates_missing_values_and_detects_association_errors() -> None:
    path = Path(__file__).resolve().parents[3] / "benchmarks" / "offer-extraction" / "run.py"
    spec = importlib.util.spec_from_file_location("offer_benchmark", path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    expected = offer().model_dump(mode="json")
    actual = dict(expected, supplier="Wrong supplier", valid_until="2028-06-30")
    scored = module.score([expected], [actual])
    assert scored["fields"]["supplier"]["populated_correct"] == 0
    assert scored["fields"]["valid_until"]["missing_correct"] == 0
    assert scored["synthetic_critical_pass"] is False
    missing = module.score([expected], [])
    assert missing["false_negatives"] == 1
    assert missing["fields"]["valid_until"]["missing_correct"] == 0
