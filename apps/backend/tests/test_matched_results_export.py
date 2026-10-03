from io import BytesIO

from openpyxl import load_workbook

from app.api.matched_results_export import build_matched_results_workbook, results_filename


def test_workbook_contains_source_fields_and_all_saved_decisions() -> None:
    summary = {
        "requestId": "IMP-123",
        "sourceFile": "Sudan_EmergencyRequest_MSF_June2024.xlsx",
        "status": "finalized",
        "partner": "MSF",
        "region": "Sudan",
        "contact": "=DANGEROUS()",
        "requestDate": "2024-06-13",
        "matchedCount": 1,
        "unmatchedCount": 1,
        "items": [
            {
                "itemId": 1, "sourceRow": 2, "requestedItemNumber": "AM500-001",
                "requested": "Amoxicillin 500mg Capsules", "quantity": 2000, "unit": "caps",
                "desiredShelfLife": "24 months", "notes": "Blister pack preferred",
                "domain": "medicine", "decision": "select_alternative",
                "product": "Amoxicillin 500mg Caps (500-ct)", "itemNumber": "AM500-CAP-500",
                "rankingScore": 94.2, "availability": "on_hand_sufficient",
                "warnings": ["Check package size"], "retrievalMethods": ["lexical", "vector"],
            },
            {
                "itemId": 2, "sourceRow": 3, "requestedItemNumber": "",
                "requested": "=HYPERLINK(\"bad\")", "quantity": 5, "unit": "packs",
                "desiredShelfLife": "", "notes": "", "domain": "medicine",
                "decision": "no_match", "product": None, "itemNumber": None,
                "rankingScore": None, "availability": None,
                "warnings": [], "retrievalMethods": [],
            },
        ],
    }

    workbook = load_workbook(BytesIO(build_matched_results_workbook(summary)))
    sheet = workbook["Matched results"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    assert sheet.max_row == 3
    assert sheet.freeze_panes == "E2"
    assert sheet.auto_filter.ref == "A1:R3"
    assert sheet.cell(2, headers["Requested item number"]).value == "AM500-001"
    assert sheet.cell(2, headers["Desired shelf life"]).value == "24 months"
    assert sheet.cell(2, headers["ERP SKU"]).value == "AM500-CAP-500"
    assert sheet.cell(2, headers["Ranking score (/100)"]).value == 94.2
    assert sheet.cell(2, headers["Decision"]).value == "Alternative selected"
    assert sheet.cell(3, headers["Decision"]).value == "Unmatched"
    assert sheet.cell(3, headers["ERP SKU"]).value is None
    assert sheet.cell(3, headers["Requested item"]).value == "'=HYPERLINK(\"bad\")"
    assert sheet.cell(3, headers["Requested item"]).data_type == "s"
    details = dict(workbook["Request details"].values)
    assert details["Source file"] == "Sudan_EmergencyRequest_MSF_June2024.xlsx"
    assert details["Contact"] == "'=DANGEROUS()"
    assert details["Unmatched lines"] == 1
    assert results_filename("IMP/123") == "matched-results-IMP-123.xlsx"


def test_workbook_links_selected_supplier_offer_without_erp_sku() -> None:
    url = "https://medeor.sharepoint.com/sites/test/offer.pdf"
    summary = {
        "requestId": "IMP-124", "sourceFile": "request.xlsx", "status": "complete",
        "partner": "", "region": "", "contact": "", "requestDate": None,
        "matchedCount": 1, "unmatchedCount": 0,
        "items": [{
            "itemId": 1, "requested": "Foley catheter CH18", "quantity": 50,
            "unit": "piece", "domain": "equipment", "decision": "accept_suggestion",
            "product": "Foley catheter CH18", "itemNumber": None,
            "candidateType": "historical_offer", "rankingScore": 100,
            "availability": "unknown", "warnings": [], "retrievalMethods": ["sharepoint_offer"],
            "provenance": [{"source_type": "sharepoint", "uri": url}],
        }],
    }
    workbook = load_workbook(BytesIO(build_matched_results_workbook(summary)))
    sheet = workbook["Matched results"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    assert sheet.cell(2, headers["ERP SKU"]).value is None
    assert sheet.cell(2, headers["Match source"]).value == "SharePoint offer"
    link = sheet.cell(2, headers["SharePoint document"])
    assert link.value == url
    assert link.hyperlink.target == url
