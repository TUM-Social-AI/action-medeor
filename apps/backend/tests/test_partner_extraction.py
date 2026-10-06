"""Requester extraction uses document/filename evidence without extra per-row calls."""

import io
import json
from unittest.mock import Mock

import pytest
from docx import Document

from app.parsing import ai_review, llm_table_classifier
from app.parsing.ai_review import PartnerReviewBatch, review_document
from app.parsing.llm_client import LlmUnavailable
from app.parsing.partner_extraction import (
    MAX_CONTEXT,
    PartnerEnvelope,
    apply_partner_result,
    basic_partner,
    capture_context,
)
from app.parsing.service import parse_upload
from app.parsing.types import ParsedDocument
from tests.test_ai_review import checked, device
from tests.test_llm_table_classifier import build_rfq_pdf, rfq_mapping
from tests.test_parsing import build_xlsx


def value(text, evidence=None, source="document"):
    return {"value": text, "evidence": evidence or text, "source": source}


@pytest.mark.parametrize(
    "filename,partner,region",
    [
        ("Anfrage 121 UHO.xlsx", "UHO", ""),
        ("Anfrage 127 Somalia UHO.xlsx", "UHO", "Somalia"),
        ("Anfrage 132 Äthiopien MWA.xlsx", "MWA", "Äthiopien"),
        ("Anfrage 141 Nigeria UHO-BVA.xlsx", "UHO-BVA", "Nigeria"),
        ("folder/Anfrage 138 Kongo BVA.xlsx", "BVA", "Kongo"),
        ("49-86 RfQ.pdf", "", ""),
        ("request.csv", "", ""),
        ("Anfrage 123 RFQ.xlsx", "", ""),
    ],
)
def test_basic_filename_patterns(filename, partner, region):
    doc = ParsedDocument()
    basic_partner(doc, filename)
    assert doc.partner == {"partner": partner, "region": region, "contact": ""}


def test_document_labels_override_filename_and_exclude_supplier_section():
    doc = ParsedDocument(
        partner_context="""action medeor labworks GmbH
Requester: Health Relief
Country | Uganda | Contact | Dr. Example
Supplier: Supplier Co
Region: Germany
Contact: Sales Person
"""
    )
    basic_partner(doc, "Anfrage 127 Somalia UHO.xlsx")
    assert doc.partner == {"partner": "Health Relief", "region": "Uganda", "contact": "Dr. Example"}


def test_conflicting_requester_labels_stay_empty_even_after_ai():
    doc = ParsedDocument(partner_context="Requester: A charity\nRequester: B charity")
    basic_partner(doc, "Anfrage 121 UHO.xlsx")
    apply_partner_result(doc, json.dumps({"partner": value("A charity")}))
    assert doc.partner["partner"] == ""
    assert doc.partner_conflicts == {"partner"}
    assert any("Conflicting" in warning for warning in doc.warnings)


def test_source_supported_ai_can_use_requester_context_outside_explicit_labels():
    doc = ParsedDocument(
        partner_context="Partial Quotations: Health Relief will accept partial quotations"
    )
    basic_partner(doc, "request.pdf")
    apply_partner_result(doc, json.dumps({"partner": value("Health Relief", doc.partner_context)}))
    assert doc.partner["partner"] == "Health Relief"
    assert doc.partner["contact"] == ""


@pytest.mark.parametrize(
    "encoded",
    [
        "not JSON",
        json.dumps({"partner": "Not an evidence object"}),
        json.dumps({"partner": value("United Health Organization", "UHO", "filename")}),
        json.dumps({"contact": value("Invented contact")}),
        json.dumps({"partner": value("UHO", source="internet")}),
    ],
)
def test_invalid_suggestions_retain_basic_values(encoded):
    doc = ParsedDocument()
    basic_partner(doc, "Anfrage 127 Somalia UHO.xlsx")
    before = dict(doc.partner)
    apply_partner_result(doc, encoded)
    assert doc.partner == before
    assert any("basic suggestions" in warning for warning in doc.warnings)


def test_supplier_and_template_evidence_are_not_requester_details():
    doc = ParsedDocument(
        partner_context="action medeor labworks GmbH\nContact: Template author\nSection A: the supplier's offer\nContact: Sales Person"
    )
    basic_partner(doc, "request.pdf")
    apply_partner_result(
        doc,
        json.dumps(
            {
                "partner": value("action medeor labworks GmbH"),
                "contact": value("Sales Person"),
            }
        ),
    )
    assert not any(doc.partner.values())


def test_explicit_requester_label_wins_over_conflicting_ai_suggestion():
    doc = ParsedDocument(partner_context="Requester: Health Relief\nA different charity")
    basic_partner(doc, "Anfrage 121 UHO.xlsx")
    apply_partner_result(doc, json.dumps({"partner": value("A different charity")}))
    assert doc.partner["partner"] == "Health Relief"


def test_metadata_is_captured_without_changing_selected_sheet_or_items():
    doc = parse_upload(
        filename="request.xlsx",
        content=build_xlsx(
            [
                ["Requester", "Health Relief"],
                ["Country", "Uganda"],
                ["Contact", "test@example.test"],
                ["Product", "Quantity", "Unit", "Type"],
                ["Diagnostic scanner", 2, "pcs", "Equipment"],
            ]
        ),
    )
    assert doc.partner == {
        "partner": "Health Relief",
        "region": "Uganda",
        "contact": "test@example.test",
    }
    assert len(doc.items) == 1
    assert doc.items[0].row == 5
    assert "Diagnostic scanner" not in doc.partner_context


def test_csv_and_word_capture_requester_metadata(monkeypatch):
    from app.parsing import llm_extractor

    monkeypatch.setattr(llm_extractor, "call_llm", Mock(side_effect=LlmUnavailable("offline")))
    csv = parse_upload(
        filename="request.csv", content=b"Requester,Health Relief\nProduct,Quantity\nScanner,2\n"
    )
    assert csv.partner["partner"] == "Health Relief"
    word = Document()
    word.add_paragraph("Requester: Health Relief")
    word.add_paragraph("Contact: test@example.test")
    content = io.BytesIO()
    word.save(content)
    doc = parse_upload(filename="request.docx", content=content.getvalue())
    assert doc.partner["partner"] == "Health Relief"
    assert doc.partner["contact"] == "test@example.test"


def test_pdf_context_includes_headings(monkeypatch):
    monkeypatch.setattr(llm_table_classifier, "call_llm", Mock(return_value=rfq_mapping()))
    doc = parse_upload(filename="Anfrage 121 UHO.pdf", content=build_rfq_pdf())
    assert len(doc.items) == 12
    assert doc.partner_context
    assert doc.partner["partner"] == "UHO"


def balanced_doc():
    doc = device()
    basic_partner(doc, "Anfrage 127 Somalia UHO.xlsx")
    doc.extraction_mode = "balanced"
    return doc


def combined_response(prompt, schema):
    response = checked(prompt, schema).model_dump()
    if schema is PartnerReviewBatch:
        response["partner_json"] = json.dumps({"partner": value("UHO", source="filename")})
    return response


def test_balanced_partner_extraction_shares_one_item_call(monkeypatch):
    doc = balanced_doc()
    call = Mock(side_effect=combined_response)
    monkeypatch.setattr(ai_review, "call_llm", call)
    review_document(doc)
    assert call.call_count == 1
    payload = json.loads(call.call_args.args[0].rsplit("\n", 1)[1])
    assert payload["partner_source"]["filename"] == doc.source_filename
    assert doc.partner["partner"] == "UHO"
    assert doc.items[0].verification_source == "ai"
    assert not doc.warnings


def test_basic_and_saved_item_reviews_do_not_extract_partner(monkeypatch):
    call = Mock(side_effect=checked)
    monkeypatch.setattr(ai_review, "call_llm", call)
    doc = balanced_doc()
    doc.extraction_mode = "basic"
    review_document(doc)
    assert "partner_source" not in json.loads(call.call_args.args[0].rsplit("\n", 1)[1])
    doc.source_filename = None
    doc.extraction_mode = "balanced"
    review_document(doc)
    assert "partner_source" not in json.loads(call.call_args.args[0].rsplit("\n", 1)[1])
    assert doc.partner["partner"] == "UHO"


@pytest.mark.parametrize(
    "partner_json", ["", "invalid", json.dumps({"contact": value("Invented")})]
)
def test_invalid_metadata_does_not_invalidate_successful_item_review(monkeypatch, partner_json):
    def call(prompt, schema):
        result = combined_response(prompt, schema)
        result["partner_json"] = partner_json
        return result

    monkeypatch.setattr(ai_review, "call_llm", call)
    doc = balanced_doc()
    review_document(doc)
    assert doc.items[0].verification_source == "ai"
    assert doc.review_summary["status"] == "completed"
    assert doc.warnings and doc.partner["partner"] == "UHO"


def test_provider_failure_preserves_basic_partner_without_extra_attempt(monkeypatch):
    call = Mock(side_effect=LlmUnavailable("private provider detail"))
    monkeypatch.setattr(ai_review, "call_llm", call)
    doc = balanced_doc()
    review_document(doc)
    assert call.call_count == 1
    assert doc.partner["partner"] == "UHO"
    assert all("private" not in warning for warning in doc.warnings)


def test_bounded_first_batch_shares_metadata_without_repeating_or_dropping_rows(monkeypatch):
    doc = balanced_doc()
    doc.partner_context = "Document heading " + "x" * 7000
    doc.items = [device().items[0] for _ in range(205)]
    seen, metadata = [], []

    def call(prompt, schema):
        payload = json.loads(prompt.rsplit("\n", 1)[1])
        seen.extend(row["row_id"] for row in payload["rows"])
        metadata.append("partner_source" in payload)
        assert len(prompt) <= ai_review.MAX_CHARACTERS
        assert len(payload["rows"]) <= 100
        return combined_response(prompt, schema)

    monkeypatch.setattr(ai_review, "call_llm", call)
    review_document(doc)
    assert len(seen) == len(set(seen)) == 205
    assert metadata[0] and sum(metadata) == 1
    assert doc.review_summary["checked"] == 205


def test_metadata_only_call_when_context_cannot_share_a_row_batch(monkeypatch):
    doc = balanced_doc()
    doc.items = device(name="Diagnostic scanner " + "x" * 11000).items
    doc.partner_context = "x" * MAX_CONTEXT
    schemas = []

    def call(prompt, schema):
        schemas.append(schema)
        assert len(prompt) <= ai_review.MAX_CHARACTERS
        if schema is PartnerEnvelope:
            return PartnerEnvelope(partner_json='{"partner":null,"region":null,"contact":null}')
        return checked(prompt, schema)

    monkeypatch.setattr(ai_review, "call_llm", call)
    review_document(doc)
    assert schemas == [PartnerEnvelope, ai_review.ReviewBatch]
    assert doc.items[0].verification_source == "ai"


def test_metadata_displaced_rows_join_later_batches_without_an_extra_call(monkeypatch):
    doc = balanced_doc()
    doc.partner_context = "x" * 4000
    doc.items = [device().items[0] for _ in range(205)]
    for index, item in enumerate(doc.items, 1):
        item.review_id = f"row-{index:04d}"
    # Ordinary batches fit 100 rows; metadata needs some of the first batch's space.
    monkeypatch.setattr(ai_review, "MAX_CHARACTERS", len(ai_review.INSTRUCTIONS) + len(ai_review._payload(doc.items[:100])) + 500)
    sizes = []
    def call(prompt, schema):
        sizes.append(len(json.loads(prompt.rsplit("\n", 1)[1])["rows"]))
        assert len(prompt) <= ai_review.MAX_CHARACTERS
        return combined_response(prompt, schema)
    monkeypatch.setattr(ai_review, "call_llm", call)
    review_document(doc)
    assert len(sizes) == 3
    assert sizes[0] < 100 and sizes[1] == 100
    assert sum(sizes) == 205


def test_metadata_only_document_gets_one_call_and_context_is_bounded(monkeypatch):
    doc = balanced_doc()
    doc.items = []
    capture_context(doc, "x" * 30_000 + "\nRequester: Health Relief\nContact: test@example.test")
    assert len(doc.partner_context) <= MAX_CONTEXT
    assert "Requester: Health Relief" in doc.partner_context
    call = Mock(return_value=PartnerEnvelope(partner_json='{"partner":null}'))
    monkeypatch.setattr(ai_review, "call_llm", call)
    review_document(doc)
    assert call.call_count == 1
    assert doc.review_summary["attempts"] == 1
