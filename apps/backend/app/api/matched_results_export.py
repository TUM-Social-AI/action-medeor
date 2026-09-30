"""Excel view of saved decisions, independent of the originally uploaded file.

The import keeps ``raw_file`` and each line's source row. A later exporter can append
these normalized result fields to the original workbook without changing decision lookup.
"""

from __future__ import annotations

import re
from io import BytesIO
from typing import Any
from urllib.parse import urlparse

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

RESULT_COLUMNS = (
    ("Line", "line", 8),
    ("Source row", "sourceRow", 12),
    ("Requested item number", "requestedItemNumber", 23),
    ("Requested item", "requested", 42),
    ("Quantity", "quantity", 13),
    ("Unit", "unit", 12),
    ("Desired shelf life", "desiredShelfLife", 20),
    ("Notes", "notes", 38),
    ("Domain", "domain", 16),
    ("Decision", "decision", 23),
    ("Matched product", "product", 44),
    ("ERP SKU", "itemNumber", 21),
    ("Ranking score (/100)", "rankingScore", 21),
    ("Availability", "availability", 24),
    ("Warnings", "warnings", 42),
    ("Retrieval methods", "retrievalMethods", 28),
    ("Match source", "sourceType", 24),
    ("SharePoint document", "sourceUrl", 60),
)

DECISION_LABELS = {
    "accept_suggestion": "Top suggestion selected",
    "select_alternative": "Alternative selected",
    "manual_match": "Manual match",
    "no_match": "Unmatched",
    "procurement_required": "Procurement required",
}


def results_filename(request_id: str) -> str:
    safe_id = re.sub(r"[^A-Za-z0-9_-]", "-", request_id).strip("-") or "request"
    return f"matched-results-{safe_id}.xlsx"


def _sharepoint_url(item: dict[str, Any]) -> str | None:
    for source in item.get("provenance") or []:
        if source.get("source_type") != "sharepoint":
            continue
        uri = source.get("uri")
        if not isinstance(uri, str):
            continue
        parsed = urlparse(uri)
        if parsed.scheme == "https" and (parsed.hostname or "").endswith(".sharepoint.com"):
            return uri
    return None


def normalized_results(summary: dict[str, Any]) -> list[dict[str, Any]]:
    """One row per requested line, including explicitly unmatched lines."""
    rows = []
    for position, item in enumerate(summary["items"], start=1):
        row = {**item, "line": position}
        row["decision"] = DECISION_LABELS.get(item["decision"], item["decision"])
        row["warnings"] = "\n".join(item["warnings"])
        row["retrievalMethods"] = ", ".join(item["retrievalMethods"])
        row["sourceType"] = (
            "SharePoint offer" if item.get("candidateType") == "historical_offer"
            else "ERP catalog" if item.get("itemNumber") else "Unmatched"
        )
        row["sourceUrl"] = _sharepoint_url(item)
        rows.append(row)
    return rows


def _excel_value(value: Any) -> Any:
    if not isinstance(value, str):
        return value
    # Excel interprets these leading characters as formulas when a workbook is opened.
    value = re.sub(r"[\x00-\x08\x0b-\x0c\x0e-\x1f]", "", value)
    if value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def build_matched_results_workbook(summary: dict[str, Any]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Matched results"
    sheet.append([label for label, _, _ in RESULT_COLUMNS])
    sheet.freeze_panes = "E2"
    sheet.sheet_view.showGridLines = False

    header_fill = PatternFill("solid", fgColor="1B4E8A")
    for column, (_, _, width) in enumerate(RESULT_COLUMNS, start=1):
        cell = sheet.cell(1, column)
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.row_dimensions[1].height = 32

    for row in normalized_results(summary):
        sheet.append([_excel_value(row.get(key)) for _, key, _ in RESULT_COLUMNS])
        row_number = sheet.max_row
        if row["sourceUrl"]:
            link_cell = sheet.cell(row_number, len(RESULT_COLUMNS))
            link_cell.hyperlink = row["sourceUrl"]
            link_cell.style = "Hyperlink"
        unmatched = row["decision"] == "Unmatched"
        for cell in sheet[row_number]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if unmatched:
                cell.fill = PatternFill("solid", fgColor="F2F4F7")
        sheet.row_dimensions[row_number].height = 32 if row["warnings"] else 25
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(RESULT_COLUMNS))}{sheet.max_row}"
    for row_number in range(2, sheet.max_row + 1):
        sheet.cell(row_number, 13).number_format = "0.0"

    details = workbook.create_sheet("Request details")
    details.sheet_view.showGridLines = False
    details.append(["Field", "Value"])
    for cell in details[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
    detail_fields = (
        ("Request ID", "requestId"),
        ("Source file", "sourceFile"),
        ("Status", "status"),
        ("Organization", "partner"),
        ("Region", "region"),
        ("Contact", "contact"),
        ("Request date", "requestDate"),
        ("Matched lines", "matchedCount"),
        ("Unmatched lines", "unmatchedCount"),
    )
    for label, key in detail_fields:
        details.append([label, _excel_value(summary.get(key))])
    details.column_dimensions["A"].width = 24
    details.column_dimensions["B"].width = 48
    details.freeze_panes = "A2"

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
