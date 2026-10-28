"""Explicit, repeatable sample import for every supported review problem.

Run from apps/backend: .venv/bin/python -m app.jobs.seed_review_example
No provider calls, migrations, or startup seeding. Running again preserves review progress.
The stored workbook and all AI findings are illustrative, clearly labelled sample data.
"""

from __future__ import annotations

import asyncio
from datetime import UTC, datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from app.db.models import ImportRequestRow
from app.db.repository import get_request_by_id, save_parsed_request
from app.db.session import async_session, engine
from app.parsing.types import ParsedDocument, ParsedLineItem
from app.review.engine import synchronize

REQUEST_ID = "EXAMPLE-REVIEW-ALL-CASES"
FILE_NAME = "EXAMPLE - All review cases.xlsx"
HEADERS = [
    "Product",
    "Requested total",
    "Unit",
    "Type",
    "Packs",
    "Units per pack",
    "Notes",
    "Example case",
    "Try this",
    "Priority",
    "Shelf life",
    "Item number",
    "Storage conditions",
]
ROLES = {
    "0": "name",
    "1": "quantity",
    "2": "unit",
    "3": "domain",
    "4": "quantity_packs",
    "5": "units_per_pack",
    "6": "notes",
    "9": "priority",
    "10": "shelf_life",
    "11": "item_number",
}


def build_review_example() -> tuple[ParsedDocument, bytes]:
    """Build frozen extraction results as well as the actual source workbook."""
    workbook = Workbook()
    workbook.active.title = "Review cases"
    workbook.create_sheet("Column check")
    for sheet in workbook:
        sheet.append(HEADERS)
    document = ParsedDocument(
        extraction_mode="balanced",
        selected_sheet="Review cases",
        selected_sheets=["Review cases", "Column check"],
        partner={
            "partner": "Example partner — sample data",
            "region": "Demo",
            "contact": "review@example.test",
        },
        warnings=[
            "Illustrative review example: extracted values and AI findings are simulated. No AI provider was called. The source workbook includes suggested actions for every case."
        ],
        review_summary={
            "status": "completed",
            "message": "Illustrative sample data, including simulated AI provenance.",
        },
        attribute_columns=["Example case", "Try this"],
        available_columns=["Storage conditions"],
    )

    def add(
        case,
        name,
        hint,
        *,
        quantity=10,
        unit="pcs",
        domain="equipment",
        sheet="Review cases",
        source=None,
        code=None,
        message=None,
        fields=None,
        **kwargs,
    ):
        table = workbook[sheet]
        values = [
            name,
            quantity,
            unit,
            domain,
            "",
            "",
            kwargs.get("notes", ""),
            case,
            hint,
            kwargs.get("priority", "medium"),
            kwargs.get("shelf_life", ""),
            kwargs.get("item_number", ""),
            "Store dry, 15–25 °C",
        ]
        for index, value in (source or {}).items():
            values[int(index)] = value
        table.append(values)
        findings = (
            [{"code": code, "message": message or hint, "fields": fields or [], "origin": "ai"}]
            if code
            else []
        )
        item = ParsedLineItem(
            name=name,
            quantity=quantity,
            unit=unit,
            domain=domain,
            confidence=95,
            status="verified",
            verification_source="ai",
            row=table.max_row,
            excerpt=" | ".join(str(value) for value in values),
            attributes={"Example case": case, "Try this": hint},
            review_findings=findings,
            review_source={
                "worksheet": sheet,
                "headers": {str(i): h for i, h in enumerate(HEADERS)},
                "roles": dict(ROLES),
                "cells": {str(i): str(v) if v is not None else "" for i, v in enumerate(values)},
            },
        )
        for key, value in kwargs.items():
            setattr(item, key, value)
        document.items.append(item)
        return item

    add(
        "Already AI checked",
        "Digital thermometer",
        "Already ready. No extra human verification is required.",
        shelf_life="12 months",
        priority="high",
        item_number="DEMO-001",
    )
    add(
        "Already human verified",
        "Sterile syringe 5 ml",
        "Already verified by a person in this example.",
        verification_source="human",
    )
    add("Missing product type", "Clinical supply A", "Choose Equipment or Medicine.", domain=None)
    add(
        "Conflicting product type",
        "Stethoscope",
        "The source states Equipment; correct the copied Medicine type.",
        domain="medicine",
        source={3: "equipment"},
    )
    add(
        "Missing quantity",
        "Examination gloves, nitrile",
        "Enter a positive quantity, for example 25.",
        quantity=None,
    )
    add(
        "Invalid quantity",
        "Sterile gauze swabs",
        "Replace zero with a positive quantity, for example 20.",
        quantity=0,
    )
    add(
        "Copied total conflict",
        "IV cannula 22G",
        "Choose the source total of 12 or explicitly resolve the conflict.",
        quantity=9,
        source={1: 12},
    )
    add(
        "Pack-total conflict",
        "Surgical masks",
        "Source total is 800; 10 packs × 100 gives 1000. Choose which total was requested.",
        quantity=800,
        source={4: 10, 5: 100},
    )
    add(
        "Unclear unit",
        "Wound dressings",
        "Set the unit to pcs. Quantity must stay 10.",
        unit="?",
        source={2: "pcs"},
    )
    add(
        "Conflicting unit",
        "Saline 0.9%, 100 ml",
        "The source says bottles; replace vials or confirm it after inspection.",
        unit="vials",
        domain="medicine",
        source={2: "bottles"},
        code="unit_conflict",
        fields=["unit"],
        message="Copied unit differs from the source unit.",
    )
    add(
        "Unclear packaging",
        "Paracetamol 500 mg tablets",
        "Inspect whether 10 counts tablets or packs, then confirm the requested unit.",
        domain="medicine",
        code="packaging_unclear",
        fields=["unit"],
        message="The requested packaging needs inspection.",
    )
    add(
        "Missing product name",
        "",
        "Recover the complete name from the source: Adult blood pressure cuff.",
        source={0: "Adult blood pressure cuff"},
    )
    add(
        "Uncertain description",
        "Infusion set ...",
        "Complete the description from the source or explicitly accept it after inspection.",
        source={0: "Infusion set, Luer lock, sterile, 20 drops/ml"},
        code="name_unclear",
        fields=["name"],
        message="The product description may be truncated.",
    )
    for name in ("Elastic bandage 8 cm", "Elastic bandage 10 cm"):
        item = add(
            "Shared column mapping",
            name,
            "Confirm the columns once for both bandage rows.",
            sheet="Column check",
            confidence=60,
            status="needs_review",
            verification_source=None,
        )
        item.review_source["mappingUncertain"] = True
    add(
        "Source coverage",
        "Sterile dressing kit",
        "Source also contains Emergency blanket, quantity 7, pcs. Add the missing item, then acknowledge reconciliation.",
        code="coverage_uncertain",
        message="A source row may be missing: Emergency blanket, quantity 7, pcs. Reconcile the source before accepting.",
    )
    add(
        "Other uncertainty",
        "Sterile surgical drape",
        "Check sterility certification. Accept after inspection or correct the note individually.",
        code="unrecognized_demo_finding",
        fields=["notes"],
        message="Sterility certification needs a manual check.",
        notes="Sterility status unconfirmed",
    )
    add(
        "Unexplained low confidence",
        "Patient transfer sheet",
        "Inspect the original row, then explicitly accept the copied values.",
        confidence=55,
        status="low_confidence",
        verification_source=None,
    )
    add(
        "Human confirmation needed",
        "Reusable kidney dish",
        "No AI verification is recorded; inspect the source and confirm.",
        confidence=None,
        status="needs_review",
        verification_source=None,
    )
    header = add(
        "Repeated header",
        "Product",
        "Remove this repeated header. You can undo removal or restore it.",
        quantity=None,
    )
    header.review_source["cells"] = {str(i): h for i, h in enumerate(HEADERS)}
    header.excerpt = " | ".join(HEADERS)
    for index, value in enumerate(HEADERS, 1):
        workbook["Review cases"].cell(header.row, index, value)
    add(
        "Overlapping problems",
        "Clinical supply B",
        "Fix the type first. Its independent pack-total conflict must remain open.",
        quantity=800,
        domain=None,
        source={4: 10, 5: 100},
    )
    add(
        "Removed item",
        "Example item removed earlier",
        "Open Removed items and restore this row. Its source and identity are retained.",
    )
    workbook["Review cases"].append(
        [
            "Emergency blanket",
            7,
            "pcs",
            "equipment",
            "",
            "",
            "Missing from extraction — add manually",
            "Coverage example",
            "Add from the item table",
        ]
    )
    document.rows_detected = len(document.items) + 1
    document.sheets = [{"name": s.title, "rows": s.max_row - 1} for s in workbook]
    document.review_summary["checked"] = len(document.items)
    for sheet in workbook:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1B4E8A")
        for column, width in {
            "A": 44,
            "B": 18,
            "C": 12,
            "D": 14,
            "E": 10,
            "F": 18,
            "G": 42,
            "H": 30,
            "I": 100,
        }.items():
            sheet.column_dimensions[column].width = width
    buffer = BytesIO()
    workbook.save(buffer)
    return document, buffer.getvalue()


async def seed_review_example(session: AsyncSession) -> tuple[ImportRequestRow, bool]:
    # A fixed identity and transaction lock prevent duplicate examples on concurrent/repeated runs.
    await session.execute(text("SELECT pg_advisory_xact_lock(hashtext(:key))"), {"key": REQUEST_ID})
    existing = await get_request_by_id(session, REQUEST_ID)
    if existing:
        if existing.source_file_name != FILE_NAME:
            raise ValueError("The example ID belongs to another request; refusing to overwrite it")
        return existing, False
    document, raw_file = build_review_example()
    draft = ImportRequestRow(
        request_id=REQUEST_ID,
        source_file_name="",
        workflow_status="draft",
        items=[],
        review_issues=[],
    )
    session.add(draft)
    await session.flush()
    request = await save_parsed_request(
        session, parsed=document, file_name=FILE_NAME, raw_file=raw_file, request_id=REQUEST_ID
    )
    # The example includes a removal that predates this review session, without fabricating an undo.
    removed = request.items[-1]
    removed.removed_at = datetime.now(UTC)
    removed.removed_by = "example setup"
    synchronize(request)
    await session.commit()
    return request, True


async def _main() -> None:
    try:
        async with async_session() as session:
            request, created = await seed_review_example(session)
            print(
                f"{request.request_id}: {'created' if created else 'already exists; progress preserved'}"
            )
            print(f"Source: {request.source_file_name}")
    finally:
        await engine.dispose()


if __name__ == "__main__":
    asyncio.run(_main())
