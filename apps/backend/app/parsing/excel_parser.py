"""Excel (.xlsx / .xls) parsing: sheet -> raw rows -> ParsedDocument via the shared table parser."""

import io

from app.parsing.table_parser import extract_request_priority_hint, parse_table_rows
from app.parsing.types import CustomColumnSpec, ParsedDocument


def parse_excel(
    content: bytes,
    filename: str,
    custom_columns: list[CustomColumnSpec] | None = None,
) -> ParsedDocument:
    if filename.lower().endswith(".xls"):
        rows = _read_xls(content)
    else:
        rows = _read_xlsx(content)

    if not rows:
        document = ParsedDocument()
        document.warnings.append("Workbook contained no rows")
        return document

    default_priority = extract_request_priority_hint(rows)
    return parse_table_rows(rows, default_priority=default_priority, custom_columns=custom_columns)


def _read_xlsx(content: bytes) -> list[list[object]]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        sheet = _pick_sheet(workbook)
        return [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def _pick_sheet(workbook):
    """Prefer the first sheet that actually has data over a blank cover/instructions sheet."""
    for sheet in workbook.worksheets:
        if sheet.max_row and sheet.max_row > 1 and sheet.max_column and sheet.max_column > 1:
            return sheet
    return workbook.worksheets[0]


def _read_xls(content: bytes) -> list[list[object]]:
    import xlrd

    workbook = xlrd.open_workbook(file_contents=content)
    sheet = workbook.sheet_by_index(0)
    for candidate in workbook.sheets():
        if candidate.nrows > 1 and candidate.ncols > 1:
            sheet = candidate
            break
    return [sheet.row_values(row_index) for row_index in range(sheet.nrows)]
