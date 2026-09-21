"""Word (.docx) parsing: always goes through the LLM free-text fallback.

Partner request .docx files carry the necessary information but not in a structure we can
reliably key column roles off of - unlike Excel or a well-tabulated PDF, there's no "try
heuristics first" step here. Paragraph and table text are flattened into one block and handed
straight to the shared free-text fallback (LLM, then a naive regex parse if no key is
configured - see free_text_parser.py).
"""

from app.parsing.free_text_parser import extract_from_free_text
from app.parsing.types import CustomColumnSpec, ParsedDocument


def parse_docx(content: bytes, custom_columns: list[CustomColumnSpec] | None = None) -> ParsedDocument:
    import io

    from docx import Document

    document = ParsedDocument()
    word_doc = Document(io.BytesIO(content))

    text_parts = [paragraph.text for paragraph in word_doc.paragraphs if paragraph.text.strip()]
    for table in word_doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            row_text = " | ".join(cell for cell in cells if cell)
            if row_text:
                text_parts.append(row_text)

    full_text = "\n".join(text_parts)
    if not full_text.strip():
        document.warnings.append("Word document contained no readable text")
        return document

    document.items.extend(extract_from_free_text(full_text, document, custom_columns=custom_columns))
    document.rows_detected = len(document.items)
    document.attribute_columns = list(
        dict.fromkeys(label for item in document.items for label in item.attributes)
    )
    if not document.items:
        document.warnings.append("No line items could be extracted from this document")
    return document
