"""Recover source context for legacy requests without repeating free-text extraction."""

import io

from app.parsing.review_context import text_review_source
from app.parsing.table_mapping import MappingSession
from app.parsing.table_parser import is_table_well_structured, parse_table_rows
from app.parsing.types import ParsedDocument


def fill_legacy_sources(
    document: ParsedDocument, filename: str, content: bytes, table_mappings: dict | None = None
) -> None:
    missing = [item for item in document.items if not item.manual and not item.review_source]
    if not missing:
        return
    extension = filename.lower().rsplit(".", 1)[-1]
    texts = {}
    table_sources = {}
    mapping = MappingSession(persisted=table_mappings or {})
    if extension == "pdf":
        import pdfplumber

        with pdfplumber.open(io.BytesIO(content)) as pdf:
            for index, page in enumerate(pdf.pages, 1):
                tables = page.find_tables()
                selected = max(tables, key=lambda t: len(t.rows)) if tables else None
                rows = selected.extract() if selected else None
                if rows and is_table_well_structured(rows):
                    top = selected.bbox[1]
                    context = page.crop((0, 0, page.width, top)).extract_text() if top > 0 else ""
                    parsed = parse_table_rows(
                        rows, page=index, context=context or "", mapping_session=mapping
                    )
                    table_sources.update(
                        {(item.page, item.row): item.review_source for item in parsed.items}
                    )
                    document.used_llm_fallback |= parsed.used_llm_fallback
                else:
                    texts[index] = page.extract_text() or ""
    elif extension == "docx":
        from docx import Document

        source = Document(io.BytesIO(content))
        texts[0] = "\n".join(
            [
                *(p.text for p in source.paragraphs),
                *(
                    " | ".join(c.text for c in row.cells)
                    for table in source.tables
                    for row in table.rows
                ),
            ]
        )
    else:
        from app.parsing.service import parse_upload

        parsed = parse_upload(filename=filename, content=content, table_mappings=table_mappings)
        table_sources = {(item.page, item.row): item.review_source for item in parsed.items}
        document.used_llm_fallback |= parsed.used_llm_fallback
    for item in missing:
        source = table_sources.get((item.page, item.row))
        if source is not None:
            item.review_source = source
        else:
            text = texts.get(item.page, texts.get(0, ""))
            item.review_source = text_review_source(text, item) if text else {"mappingUncertain": True}
