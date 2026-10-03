"""Reproducible PDF fixtures and independently annotated workbook expectations.

No model or production extractor is used to prepare labels. Workbook row selections and
date/reference annotations below were reviewed against the supplied source cells.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import random
import re
from decimal import Decimal
from pathlib import Path
from xml.sax.saxutils import escape

from openpyxl import load_workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

ROOT = Path(__file__).resolve().parents[2]
DEFAULT_DATA = ROOT / "data" / "anonymized data"

# Explicit included source rows, supplier column. Quantity-only/enquiry/internal-stock rows
# are excluded. A product title in the status cell is legitimate offered-description evidence.
ROWS = {
    "Anfrage 121 UHO.xlsx": {"G": list(range(5, 32))},
    "Anfrage 127 Somalia UHO.xlsx": {
        "H": [
            8,
            9,
            10,
            12,
            15,
            16,
            17,
            18,
            19,
            22,
            23,
            24,
            25,
            26,
            29,
            30,
            31,
            32,
            33,
            34,
            35,
        ],
        "N": [13, 15, 16, 17, 23, 29, 30, 31],
    },
    "Anfrage 132 Äthiopien MWA.xlsx": {
        "G": [
            7,
            8,
            9,
            10,
            11,
            12,
            14,
            15,
            16,
            17,
            18,
            20,
            21,
            23,
            26,
            29,
            30,
            31,
            32,
            33,
            34,
            35,
            36,
            37,
            39,
            41,
            42,
            48,
            49,
            50,
            51,
            52,
            53,
            55,
            56,
            57,
            59,
        ],
    },
    "Anfrage 138 Kongo BVA.xlsx": {
        "G": [7, 8, 10, 12, 13, 14, 15, 16, 17, 18, 19, 20, 21, 22, 23],
        "M": [18, 23],
    },
    "Anfrage 141 Nigeria UHO-BVA.xlsx": {
        "G": [
            7,
            8,
            9,
            10,
            11,
            12,
            14,
            15,
            16,
            17,
            18,
            19,
            20,
            21,
            22,
            23,
            24,
            27,
            28,
            29,
            30,
            31,
            32,
            35,
            38,
            40,
            41,
            43,
            44,
            45,
            46,
            47,
            50,
            52,
            53,
            57,
            58,
            59,
            60,
            61,
            62,
            63,
            64,
            65,
            66,
            67,
            68,
            70,
            71,
            72,
            73,
            74,
            75,
            76,
            77,
            78,
            79,
            80,
            81,
            82,
            83,
            85,
            86,
            88,
            89,
            90,
            91,
            93,
            95,
        ],
    },
}

# Explicit status annotations. Missing/partial dates deliberately have no ISO date.
STATUS = {
    "Angebot 25028186": ("25028186", None, None),
    "Angebot per Mail 06.03.2026": (None, "2026-03-06", "sent"),
    "Angebot 103568Nov. 25": ("103568", None, None),
    "SQ-025002 v. 10.03.2026": ("SQ-025002", "2026-03-10", "issued"),
    "53877698 v. 11.03.2026": ("53877698", "2026-03-11", "issued"),
    "213351 v. 10.03.2026": ("213351", "2026-03-10", "issued"),
    "Angebot 122323 v. 09.04.2026": ("122323", "2026-04-09", "issued"),
    "per Mail 14.04.2026": (None, "2026-04-14", "sent"),
    "Angebot 105091 v. 09.04.26": ("105091", "2026-04-09", "issued"),
    "PI 267 v. 09.04.2026": ("267", "2026-04-09", "issued"),
    "AN32612195 v. 13.04.2026": ("AN32612195", "2026-04-13", "issued"),
    "Angebot 214198 v. 10.04.2026": ("214198", "2026-04-10", "issued"),
    "DE064438 v. 12.04.2026": ("DE064438", "2026-04-12", "issued"),
    "Angebot 5423158 v. 24.04.2026": ("5423158", "2026-04-24", "issued"),
    "Stocklist v. 20.04.2026": (None, "2026-04-20", "source"),
    "Angebot 5414340 v. 14.04.2026": ("5414340", "2026-04-14", "issued"),
    "SQ-023745": ("SQ-023745", None, None),
    "Angebot SQ-025355 v. 22.04.2026": ("SQ-025355", "2026-04-22", "issued"),
    "122345 v. 16.04.2026": ("122345", "2026-04-16", "issued"),
    "122378 v. 22.04.2026": ("122378", "2026-04-22", "issued"),
    "Angebot 1322 vom 22.12.2025": ("1322", "2025-12-22", "issued"),
    "per Mail am 21.04.2026": (None, "2026-04-21", "sent"),
    "Angebot 105239 v. 21.04.2026": ("105239", "2026-04-21", "issued"),
    "553383": ("553383", None, None),
    "Angebot 214407 vpm 20.04.2026": ("214407", "2026-04-20", "issued"),
    "2026-500001": ("2026-500001", None, None),
    "per Mail 06.08.2025": (None, "2025-08-06", "sent"),
    "AN32612213 v. 30.04.2026": ("AN32612213", "2026-04-30", "issued"),
    "5427802 v. 04.05.2026": ("5427802", "2026-05-04", "issued"),
    "Angebot 5422363 v. 23.04.2026": ("5422363", "2026-04-23", "issued"),
    "5387832 Angebot Puderbach": ("5387832", None, None),
    "AN32612222 v. 11.05.2026": ("AN32612222", "2026-05-11", "issued"),
    "Angebot per Mail 08.05.2026": (None, "2026-05-08", "sent"),
    "11428900 v. 18.08.2025 + 25 %": ("11428900", "2025-08-18", "issued"),
    "Angebot 5434224 v. 11.05.2026": ("5434224", "2026-05-11", "issued"),
    "Angebot 5423158 Centramed": ("5423158", None, None),
    "26-051662 vom 07.05.2026": ("26-051662", "2026-05-07", "issued"),
    "offer 25028579 v. 11.05.2026": ("25028579", "2026-05-11", "issued"),
    "per Mail 06.05.2026": (None, "2026-05-06", "sent"),
    "per Mail 12.4. 2025 + 10%": (None, "2025-04-12", "sent"),
    "Angebot DE064538 v. 07.05.2026": ("DE064538", "2026-05-07", "issued"),
    "SQ-025355 v. 22.04.2026": ("SQ-025355", "2026-04-22", "issued"),
    "Mal Fr. Lelleik 08.05.2026": (None, "2026-05-08", "sent"),
}


def base_offer(source_id: str, supplier: str | None, description: str) -> dict:
    return {
        "source_id": source_id,
        "alternative_index": 1,
        "supplier": supplier,
        "item_description": description,
        "offer_reference": None,
        "offer_date": None,
        "date_kind": None,
        "valid_until": None,
        "validity_text": None,
        "price_text": None,
        "price_amount": None,
        "currency": None,
        "price_basis": None,
        "price_conflict": False,
    }


def workbook_labels(data: Path) -> dict:
    result = {}
    for name, blocks in ROWS.items():
        path = data / name
        book = load_workbook(path, data_only=True)
        sheet = book.worksheets[0]
        expected = []
        for supplier_col, rows in blocks.items():
            start = sheet[f"{supplier_col}1"].column
            is_first = supplier_col in ("G", "H")
            status_col = start + (4 if is_first else 6)
            for row in rows:
                supplier_raw = sheet.cell(row, start).value
                supplier = re.sub(
                    r"\s+angefragt\s*$", "", str(supplier_raw or ""), flags=re.I
                )
                supplier = re.sub(r"\s+\d{7}$", "", supplier.strip()) or None
                description = sheet.cell(row, start + 1).value
                status = str(sheet.cell(row, status_col).value or "").strip()
                if not description or not str(description).strip():
                    # Individually reviewed product titles in status cells; otherwise reference
                    # evidence permits keeping the request identity with an explicit warning.
                    if (
                        name == "Anfrage 127 Somalia UHO.xlsx"
                        and row == 15
                        or name == "Anfrage 132 Äthiopien MWA.xlsx"
                        and row in (39, 41, 42)
                    ):
                        description = status
                    else:
                        description = sheet.cell(row, 2).value
                offer = base_offer(
                    f"{sheet.title}!{supplier_col}{row}",
                    supplier,
                    str(description).strip(),
                )
                annotation = STATUS.get(status.split("\n")[0].strip())
                if annotation:
                    (
                        offer["offer_reference"],
                        offer["offer_date"],
                        offer["date_kind"],
                    ) = annotation
                elif re.search(r"\d{1,2}\.\d{1,2}\.", status):
                    raise ValueError(
                        f"Unannotated dated status: {name} row {row}: {status}"
                    )
                price_col = start + 4 if not is_first else None
                price_raw = sheet.cell(row, price_col).value if price_col else None
                if name == "Anfrage 141 Nigeria UHO-BVA.xlsx" and row == 60:
                    price_raw = sheet.cell(row, start + 3).value
                variants = [offer]
                if (
                    name == "Anfrage 127 Somalia UHO.xlsx"
                    and supplier_col == "N"
                    and row == 13
                ):
                    variants = []
                    for i, (desc, price) in enumerate(
                        zip(
                            str(description).split("\n"),
                            str(price_raw).split("\n\n"),
                            strict=True,
                        ),
                        1,
                    ):
                        variant = dict(
                            offer,
                            alternative_index=i,
                            item_description=desc.removeprefix("Alternative: "),
                        )
                        variants.append(variant)
                        apply_price(variant, price)
                elif price_raw:
                    apply_price(offer, str(price_raw))
                if (
                    name == "Anfrage 138 Kongo BVA.xlsx"
                    and supplier_col == "M"
                    and row == 18
                ):
                    offer.update(price_amount=None, price_conflict=True)
                for variant in variants:
                    variant["source_evidence"] = {
                        c.coordinate: str(c.value).strip()
                        for c in sheet[row][start - 1 : status_col + 1]
                        if c.value is not None
                    }
                    expected.append(variant)
        result[name] = {
            "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            "synthetic": False,
            "offers": expected,
        }
        book.close()
    return result


def apply_price(offer: dict, text: str) -> None:
    amount, basis = text.split(" / ")
    offer.update(
        price_text=text,
        price_amount=amount.replace("€", "").replace(",", "."),
        currency="EUR",
        price_basis=basis,
    )


def pdf_labels(data: Path) -> dict:
    pdfmetrics.registerFont(
        TTFont("FixtureSans", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")
    )
    styles = getSampleStyleSheet()
    for style in styles.byName.values():
        style.fontName = "FixtureSans"
    rng = random.Random(20261002)
    suppliers = [
        "Aster Medical GmbH",
        "Boreal Instruments Ltd",
        "Cedar Klinikbedarf GmbH",
        "Delta Care Supplies Ltd",
        "Elara Pharma GmbH",
        "Fir Medicines Ltd",
        "Grove Arzneimittel GmbH",
        "Harbor Therapeutics Ltd",
    ]
    products = [
        ["Sterile needle 19G 1.1 x 40 mm", "Disposable syringe Luer 5 ml"],
        ["Foley catheter CH18 2-way 40 cm", "Infusion stand on castors"],
        ["Mayo-Hegar Nadelhalter 18 cm", "Halsted-Mosquito Klemme gebogen 12,5 cm"],
        ["Examination lamp LED mobile", "Foldable wheelchair adult"],
        ["Amoxicillin 500 mg capsules", "Paracetamol 500 mg tablets"],
        [
            "Ceftriaxone 1 g powder for injection vial",
            "Oral rehydration salts 20.5 g sachet",
        ],
        ["Metformin 500 mg Tabletten", "Ibuprofen 200 mg Tabletten"],
        [
            "Insulin human 100 IU/ml 10 ml vial",
            "Salbutamol 100 mcg/dose inhaler 200 doses",
        ],
    ]
    bases = [
        ["100 pieces", "piece"],
        ["box", "piece"],
        ["Stück", "4 Stück"],
        ["piece", "piece"],
        ["100 capsules", "100 tablets"],
        ["vial", "20 sachets"],
        ["Packung", "100 Tabletten"],
        ["vial", "inhaler"],
    ]
    result = {}
    for n, supplier in enumerate(suppliers):
        german = n in (0, 2, 4, 6)
        name = (
            f"synthetic_offer_{n + 1:02d}_{'equipment' if n < 4 else 'medication'}.pdf"
        )
        ref = f"SYN-Q-{rng.randrange(10000, 99999)}"
        offer_date = None if n == 3 else f"2026-09-{n + 10:02d}"
        date_kind = None if offer_date is None else "sent" if n == 7 else "issued"
        valid_until = "2026-11-30" if n in (0, 1, 4, 7) else None
        validity = (
            "Gültig bis: 30.11.2026" if german else "Offer valid until: 2026-11-30"
        )
        if n in (2, 5):
            validity = (
                "Angebot gültig 30 Tage ab Angebotsdatum"
                if german
                else "Offer valid for 30 days from issue"
            )
        elif valid_until is None:
            validity = None
        expected = []
        story = []
        subtotal = Decimal("0")

        def p(text: str, style: str = "Normal") -> Paragraph:
            return Paragraph(escape(text).replace("\n", "<br/>"), styles[style])

        story.extend(
            [
                p(supplier, "Title"),
                p("Fictional supplier / synthetischer Lieferant"),
                p("42 Example Avenue, 12345 Sampletown · sales@example.invalid"),
                Spacer(1, 20),
                p("ANGEBOT" if german else "QUOTATION / PRO FORMA", "Heading1"),
                p(f"{'Angebotsnummer' if german else 'Quotation number'}: {ref}"),
                p("Customer: Humanitarian Procurement Example"),
                p("Request: SAMPLE-REQ-2026"),
            ]
        )
        if offer_date:
            printed_date = (
                ".".join(reversed(offer_date.split("-"))) if german else offer_date
            )
            label = "Sent on" if n == 7 else "Angebotsdatum" if german else "Issue date"
            story.append(p(f"{label}: {printed_date}"))
        if validity:
            story.append(p(validity))
        if n >= 4:
            story.append(
                p(
                    "Product expiry / EXP: 2028-06-30. Minimum requested shelf life: 18 months."
                )
            )
        story.extend(
            [
                p("Delivery: 2-3 weeks after order. Prices exclude VAT and freight."),
                Spacer(1, 18),
            ]
        )
        for i, description in enumerate(products[n]):
            page = 2 if n == 7 and i == 1 else 1
            if page == 2:
                story.extend(
                    [
                        PageBreak(),
                        p(f"{supplier} — {ref}", "Heading2"),
                        p("Quotation continued"),
                        Spacer(1, 15),
                    ]
                )
            amount = f"{rng.randrange(200, 80000) / 100:.2f}"
            printed_amount = amount.replace(".", ",") if german else amount
            basis = bases[n][i]
            price = f"{printed_amount} {'€' if german else 'EUR'} / {basis}"
            quantity = rng.randrange(2, 16)
            denominator = int(basis.split()[0]) if basis.split()[0].isdigit() else 1
            quantity *= denominator
            subtotal += Decimal(amount) * quantity / denominator
            # Include headers and totals around real quotation rows; line indices restart per page.
            item_row = [
                [
                    p("Pos"),
                    p("Artikel / Description"),
                    p("Menge / Qty"),
                    p("Preis / Quoted price"),
                ],
                [p(str(i + 1)), p(description), p(str(quantity)), p(price)],
            ]
            table = Table(item_row, colWidths=[40, 230, 65, 155], hAlign="LEFT")
            table.setStyle(
                TableStyle(
                    [
                        (
                            "BACKGROUND",
                            (0, 0),
                            (-1, 0),
                            colors.HexColor(["#dae7ef", "#eee8dc", "#e0eddd"][n % 3]),
                        ),
                        ("BOX", (0, 0), (-1, -1), 0.6, colors.grey),
                        ("INNERGRID", (0, 0), (-1, -1), 0.3, colors.lightgrey),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 12),
                    ]
                )
            )
            story.extend([table, Spacer(1, 12)])
            offer = base_offer(
                f"page:{page}:line:{1 if page == 2 else i + 1}", supplier, description
            )
            offer.update(
                offer_reference=ref,
                offer_date=offer_date,
                date_kind=date_kind,
                valid_until=valid_until,
                validity_text=validity,
                price_text=price,
                price_amount=amount,
                currency="EUR",
                price_basis=basis,
            )
            expected.append(offer)
        freight = Decimal("85.00")
        vat = ((subtotal + freight) * Decimal("0.19")).quantize(Decimal("0.01"))
        total = subtotal + freight + vat
        story.extend(
            [
                Spacer(1, 25),
                p(f"Subtotal: EUR {subtotal:,.2f}"),
                p("Freight: EUR 85.00"),
                p(f"VAT 19%: EUR {vat:,.2f}"),
                p(f"Grand total: EUR {total:,.2f}"),
                p(
                    "Minimum order value: EUR 1,000.00. Payment: bank transfer in advance."
                ),
                p("SYNTHETIC QUOTATION — FOR EXTRACTION TESTING ONLY."),
            ]
        )

        def footer(canvas, doc):
            canvas.setFont("FixtureSans", 8)
            canvas.drawString(42, 28, f"{supplier} | {ref} | page {doc.page}")

        path = data / name
        SimpleDocTemplate(
            str(path),
            pagesize=A4,
            leftMargin=42,
            rightMargin=42,
            topMargin=40,
            bottomMargin=45,
            invariant=1,
        ).build(story, onFirstPage=footer, onLaterPages=footer)
        result[name] = {
            "synthetic": True,
            "offers": expected,
            "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        }
    return result


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--data-dir", type=Path, default=DEFAULT_DATA)
    parser.add_argument(
        "--labels", type=Path, default=DEFAULT_DATA / "offer-extraction-labels.json"
    )
    args = parser.parse_args()
    args.data_dir.mkdir(parents=True, exist_ok=True)
    documents = workbook_labels(args.data_dir)
    documents.update(pdf_labels(args.data_dir))
    labels = {
        "label_version": "source-reviewed-v2",
        "seed": 20261002,
        "method": "Explicit source row selections and manual date/reference annotations; no LLM",
        "documents": documents,
    }
    args.labels.parent.mkdir(parents=True, exist_ok=True)
    args.labels.write_text(json.dumps(labels, ensure_ascii=False, indent=2) + "\n")
    print(
        f"Created eight PDFs and labels for {len(documents)} documents at {args.labels}"
    )
    print({name: len(doc["offers"]) for name, doc in documents.items()})


if __name__ == "__main__":
    main()
